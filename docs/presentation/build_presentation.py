"""Build the German simulation presentation from the inspected repository.

Run from the repository root: .venv/bin/python docs/presentation/build_presentation.py
Requires python-pptx, Pillow and ReportLab, only for this document tool.
"""
from __future__ import annotations

import ast
import hashlib
import json
from pathlib import Path
import subprocess

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from reportlab.pdfgen.canvas import Canvas as PDFCanvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
W, H = 1600, 900
BG, PANEL, INK, MUTED = '#101F2E', '#1B3042', '#F3F5F3', '#B6C6CC'
TEAL, GOLD, RED = '#54D1BA', '#F5C66C', '#EF9582'
FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
SLIDES = []


def add(section, title, subtitle, cards, takeaway, sources, notes, kind='cards'):
    SLIDES.append(dict(section=section, title=title, subtitle=subtitle,
                       cards=cards, takeaway=takeaway, sources=sources,
                       notes=notes, kind=kind))


add('SYSTEMÜBERBLICK', 'U-Jagd: Die Simulation verstehen',
    'Wie Welt, Sensoren, Entscheidungen und Wirkung ineinandergreifen',
    [('1.1.0', 'Codebasis und lokaler Arbeitsstand'),
     ('9 Stationen', 'Ein gemeinsames Schiff und Lagebild'),
     ('24.09.2026', 'Analyse des vorliegenden Repositorys')],
    'Technische Präsentation mit Abläufen, Modellen, Abhängigkeiten und Grenzen.',
    ['src/core/version.py', 'src/core/game.py', 'pyproject.toml'],
    'Diese Präsentation beschreibt die ausführbare Implementierung im lokalen Arbeitsbaum. Dieser enthält bereits Änderungen gegenüber dem letzten Commit, unter anderem an Wetter, Sonar und Browseroberfläche. Versionsnummer, Speicherformat und Protokoll sind getrennte Verträge: Anwendung 1.1.0, Save v12 und Remote Crew v2. Die Folien dienen als Architektur- und Simulationsführung. Physikalische Formeln erläutern die implementierten Näherungen; sie sind keine Aussage über nachgewiesene Leistungsdaten realer Marinesysteme. Screenshots sind vorhandene Repository-Abbildungen und können vom aktuellen Bildschirm abweichen.', 'cover')

add('SYSTEMÜBERBLICK', 'Der rote Faden', 'Vom Simulationskern bis zum gemeinsamen Einsatz',
    [('01 · Grundlagen', 'Architektur, Zeit, Einheiten, Daten und Welt'),
     ('02 · Wahrnehmung', 'Akustik, Sonar, TMA, Radar, Funk und Fusion'),
     ('03 · Wechselwirkung', 'KI, Waffen, Schaden, Crew, Speicherung und Qualität')],
    'Leitfrage jeder Folie: Was geht hinein, was wird berechnet, wer verwendet das Ergebnis?',
    ['src/core/game.py'],
    'Die Reihenfolge folgt einer Ursache-Wirkungs-Kette. Zuerst entsteht eine Welt mit Einheiten und Umweltzustand. Sensoren erzeugen daraus unvollständige Beobachtungen. Menschen und KI entscheiden anhand dieser Daten. Stellbefehle und Waffen verändern wiederum die Welt. Querschnittssysteme wie Speicherung, Netzwerk und Darstellung halten diesen Prozess bedienbar und reproduzierbar. Die Präsentation ist ausführlich angelegt; für einen kürzeren Vortrag lassen sich die Formel- und Implementierungsfolien überspringen.')

add('SYSTEMÜBERBLICK', 'Eine Welt, mehrere Wissensstände', 'Das zentrale Prinzip der gesamten Simulation',
    [('Weltzustand', 'Position, Tiefe, Kurs, Energie, Schaden und Emissionen der Einheiten.'),
     ('Messung', 'Ausbreitung, Abschattung, Störungen und Messfehler begrenzen die Wahrnehmung.'),
     ('Beobachtung', 'Kontakte und Tracks enthalten Quelle, Alter, Qualität und gegebenenfalls einen Fix.'),
     ('Handlung', 'Bediener, KI und Waffen nutzen die jeweils verfügbaren Informationen.')],
    'Handlungen wirken auf den Weltzustand zurück – der Informationskreislauf beginnt erneut.',
    ['src/core/game.py', 'src/sensors/tracks.py', 'src/sensors/platform.py', 'src/ui/observations.py'],
    'Physik und Sensorerzeugung dürfen reale Simulationspositionen lesen: Nur so lassen sich Ausbreitung und Kollisionen bestimmen. Die daraus erzeugte Beobachtung ist jedoch ein eigenes Datenprodukt. Eine passive Peilung ist kein Entfernungsfix. Bedienoberflächen und Fernclients sollen ausschließlich freigegebene Beobachtungen erhalten. Eigenschiff, eigene beauftragte Mittel und bekannte Geografie sind legitime Ausnahmen. Terminale Waffensuchköpfe dürfen intern Kandidaten prüfen. SimLog ist eine ausdrücklich freigeschaltete Diagnoseausnahme mit abgetrennten Zustandsschnappschüssen.', 'flow')

add('SYSTEMÜBERBLICK', 'Die Architektur in sieben Bausteinen', 'Game verbindet die Fachmodule und legt die Reihenfolge fest',
    [('Daten & Konfiguration', 'JSON-Kataloge · Loadouts · Szenarien'),
     ('Welt & Physik', 'Umwelt · Bewegung · Gelände'),
     ('Einheiten & Waffen', 'Schiff · Gegner · Luft · Wirkmittel'),
     ('Sensoren & Tracks', 'Sonar · Radar · ESM · TMA'),
     ('Game', 'Initialisierung · Update · Befehle · Mission'),
     ('Bedienung & Netzwerk', 'Stationen · Autocrew · Remote Crew'),
     ('Persistenz & Ausgabe', 'Save v12 · UI · Audio · SimLog')],
    'Fachmodelle berechnen lokale Ergebnisse; der zentrale Ablauf macht daraus einen zusammenhängenden Einsatz.',
    ['src/core/game.py', 'src/physics/ship_dynamics.py', 'src/commander/local.py'],
    'Die Darstellung zeigt konzeptionelle Laufzeitabhängigkeiten, keinen vollständigen Importgraphen. Game ist der umfangreiche Integrationspunkt: Es baut Einheiten auf, ruft Modelle in definierter Reihenfolge auf, überträgt Beobachtungen und verarbeitet Wirkungen. Reine physikalische Module liefern weitgehend zustandslose Funktionen. Zustände gehören Welt, Schiffen, Waffen, Sensoren und Missionssteuerung. Die beigefügte Quelleninventur erfasst zusätzlich die Python-Module und ihre direkten internen Imports.', 'architecture')

add('SYSTEMÜBERBLICK', 'Daten legen Fähigkeiten fest', 'Katalogwerte werden in Laufzeitmodelle übersetzt',
    [('Kontaktkatalog', 'Profile für Plattformen, Akustik, Torpedos, Täuschkörper und Emittenten.'),
     ('Eigene Ausrüstung', 'Loadouts definieren Magazine, Rohre, Abwehrmittel und Rumpfparameter.'),
     ('Laufzeitbindung', 'Loader und Validatoren prüfen Daten; Entitäten erhalten konkrete Profile und Zustände.')],
    'Ein Datenfeld wirkt erst dann auf das Spiel, wenn Laufzeitcode es tatsächlich verwendet.',
    ['src/data/catalog.py', 'data/contacts/subs.json', 'data/loadouts/ownship.json', 'data/loadouts/ownship_hull.json'],
    'Kataloge sind die zentrale Quelle für eingebaute Plattform- und Signaturprofile. Der Loader macht aus JSON strukturierte Profile und überprüft Referenzen. Einheiten verbinden diese Profile mit individuellen Zuständen: etwa verbleibender Munition, Batteriespeicher oder Sensorzyklen. Rumpfparameter des Eigenschiffs liegen in einem separaten Loadout. Das ist von editorseitig akzeptierten Feldern zu unterscheiden: Eine gültige Definition ist noch keine Zusicherung, dass jedes Feld von der Simulation unterstützt wird. Laufzeitsnapshots der relevanten Katalogdaten gehören zur Speicherfortsetzung.')

add('ZEIT & ABLAUF', 'Zeit und Koordinaten sind gemeinsame Verträge', 'Der aktuelle Arbeitsstand läuft ausschließlich mit Faktor 1',
    [('Raum', 'Position und Entfernung: NM\nGeschwindigkeit: kn\nTiefe: m\nFrequenz: Hz'),
     ('Orientierung', '0° = Norden, 90° = Osten\nx wächst nach Osten\ny wächst nach Süden'),
     ('Zeit', '1 s Simulation = 1 s bei normalem Lauf\nUhrzeit und Mission folgen dt\nPause stoppt die Kernsimulation')],
    'Weg pro Sekunde = Geschwindigkeit in kn / 3600. Beispiel: 18 kn → 0,005 NM/s.',
    ['src/core/config.py', 'src/core/game.py', 'src/world/world.py'],
    'Aktuell enthält TIME_SCALE_STEPS nur den Wert 1; Game.time_scale liefert konstant 1. Ältere Zeitrafferbeschreibungen sind daher kein aktuelles Bedienversprechen. GAME_TIME_PER_SEC ist 1/60 Spielminute pro Simulationssekunde; World.update dividiert noch einmal durch 60 zur Umrechnung in Stunden. Die Weltuhr läuft somit ebenfalls in Echtzeit. Bewegung folgt x += sin(Kurs) * Strecke und y -= cos(Kurs) * Strecke. Bei langen Frames begrenzt der Hauptloop dt; dadurch wird versäumte Wandzeit nicht vollständig nachgeholt.')

add('ZEIT & ABLAUF', 'Ein Frame: Eingaben vor der Simulation', 'Die Reihenfolge verhindert unkontrollierte Zustandsänderungen',
    [('Zeit erfassen', 'Ziel: 60 FPS. Wall-dt wird für die Simulation auf höchstens 0,1 s begrenzt.'),
     ('Befehle anwenden', 'Lokale Events, Remote-Crew-Pump und optionaler Live-Verkehr werden verarbeitet.'),
     ('Welt fortschreiben', 'Game.update prüft Sperrzustände und zerlegt dt in Physikschritte.'),
     ('Ausgabe', 'Audio aktualisieren, Station zeichnen und auf das Fenster abbilden.')],
    'Browsertransport und Rendering sind keine eigenständigen Simulationsinstanzen.',
    ['src/core/game.py', 'src/commander/local.py', 'src/network/live_traffic.py'],
    'Game.run ist die äußere Schleife. Nach den Events werden zuerst Remote-Befehle und anschließend Live-Feed-Meldungen gepumpt. Erst danach folgt Game.update. Die Kernsimulation läuft bei Pause, Menü, Spielende und blockierenden Verwaltungszuständen nicht weiter; es gibt explizite Ausnahmen für bestimmte Crew-Overlays. Wichtig ist die Platzierung des LiveTraffic-Pump vor diesem Gate: Externe Feedpflege ist nicht automatisch durch die Physikpause angehalten. Im Web-Host-Modus entfällt die native Bildausgabe.', 'flow')

add('ZEIT & ABLAUF', 'Innerhalb eines Physikschritts', 'Diese Update-Reihenfolge ist selbst eine Abhängigkeit',
    [('1 · Basis', 'Simulationszeit → Navigation/Welt → Vorräte → Sonarmechanik → Radarschwenk'),
     ('2 · Akteure', 'NPC-Sensoren bei Fälligkeit → Unterwasser-/See-Einheiten → Helikopter/Bojen'),
     ('3 · Kampf', 'Raider → Luftabwehr → Feindtorpedos → eigene Torpedos → ASROC'),
     ('4 · Lage & Folgen', 'Sonar → ESM/Funk → Schaden/Mission → Autocrew → SimLog')],
    'Autocrew liest die gerade publizierten Beobachtungen; ihre Stellbefehle wirken ab dem folgenden Schritt.',
    ['src/core/game.py'],
    'Game._update_sim erhöht zunächst sim_t und mission_time. Navigation überträgt Schadensfolgen, aktualisiert die Welt und bewegt das Eigenschiff. NPC-Sensoren laufen vor den anschließenden Akteursentscheidungen. Die Waffenreihenfolge ist ebenfalls bewusst gewählt: Ein neu eingewassertes ASROC-Payload bewegt sich erst im nächsten Substep, weil der aktuelle Schritt schon als Flugzeit verbraucht wurde. Sonar, ESM, Funk und Schaden werden über Akkumulatoren getaktet. Die dargestellten Gruppen fassen benachbarte Aufrufe zusammen; HQ-Wetterbulletins und mehrere Nebenereignisse liegen dazwischen.', 'flow')

add('ZEIT & ABLAUF', 'Mehrere Takte teilen sich eine Zeitbasis', 'Nicht jedes System braucht eine neue Rechnung pro Bild',
    [('Physik', 'Substeps bis 0,05 s', 'Bewegung, Waffen, Mechanik'),
     ('Sensorlage', 'ab 0,25 s Akkumulator', 'NPC-Sensoren und eigenes Sonar'),
     ('ESM / Funk', 'ab 0,5 s Akkumulator', 'Publikation der jeweiligen Lage'),
     ('Schaden / Mission', 'ab 0,5 s Akkumulator', 'Flutung, Brand, Zielprüfung'),
     ('Autocrew / Audio', 'stations- bzw. blockweise', 'Crew: 0,5–5 s; Empfänger: 0,25 s')],
    'Takte sind Implementierungsintervalle; die Ausführung erfolgt beim nächsten passenden Substep.',
    ['src/core/game.py', 'src/core/autocrew.py', 'src/audio/receiver.py'],
    'Die Substep-Anzahl ist auf 240 begrenzt. Im normalen Hauptloop mit dt ≤ 0,1 s sind höchstens zwei 0,05-s-Schritte notwendig; bei großen direkten update-Aufrufen ist die Obergrenze zu beachten. Publikationsintervalle sind nicht gleich Sensorbeleuchtung: Das Radar erhält beispielsweise nur dann einen neuen Look, wenn der rotierende Strahl die Peilung passiert. Der akustische Empfänger erzeugt 250-ms-Blöcke. Audioausgabe verwendet zusätzlich Wandzeit und Pufferzustände, ohne die Bewegung der Einheiten davon abhängig zu machen.', 'table')

add('WELT & UMWELT', 'Geografie ist mehr als ein Kartenhintergrund', 'Bekannte Küsten, synthetischer Meeresboden und physische Abfragen',
    [('Sektor', '128 validierte reale 500-NM-Sektoren; seed % 128 wählt im Seed-Modus den Abschnitt.'),
     ('Tiefenmodell', 'Kartentiefe + Gezeit; Unterwasserfelsen können die effektive Tiefe begrenzen.'),
     ('Abfragen', 'Landkontakt, Rumpffreiheit, Sichtlinien und Unterwasserabschattung begrenzen Aktionen.')],
    'Reale geografische Namen bedeuten keine real vermessene Bathymetrie oder reale militärische Lage.',
    ['src/world/real_coast.py', 'src/world/coastline.py', 'src/world/world.py', 'src/world/grounding.py'],
    'Zusätzlich existiert eine stilisierte feste Referenzkarte und eine gezielte Auswahl realer Sektoren. Der Seed bestimmt im entsprechenden Modus stabil den Sektor und weitere Weltmerkmale. Die Küstenartefakte sind aus festgelegten Quellen erzeugt; der Meeresboden bleibt synthetisch. Laufzeitabfragen unterscheiden kartierte Tiefe, tatsächliche Tiefe und physische Geländetiefe. Grundberührung nutzt einen Rumpfumriss, nicht nur einen Mittelpunkt. Sonar- und Sichtlinienabfragen verhindern Beobachtungen durch entsprechende Geländeabschattung.')

add('WELT & UMWELT', 'Die Umwelt verzweigt in fast alle Systeme', 'Gemeinsame Ursachen erzeugen mehrere taktische Folgen',
    [('Wind & Seegang', 'Zusatzwiderstand, Rollen/Stampfen, Eigen- und Umgebungsgeräusch, Radarclutter, Fluggrenzen.'),
     ('Regen & Sicht', 'Radarabschwächung, optische Reichweite, akustischer Hintergrund und Flugwetter.'),
     ('Ozean & Tiefe', 'Schallprofil, Reflexion, Gezeit, Bodenfreiheit, Tauchgrenzen und Drift.')],
    'Ein Wetterwechsel verändert Mobilität, Aufklärung und verfügbare Einsatzmittel gleichzeitig.',
    ['src/world/world.py', 'src/world/ocean.py', 'src/physics/ship_dynamics.py', 'src/sensors/radar.py'],
    'Die Wetterendpunkte werden deterministisch abgeleitet und weich ineinander überführt. Die Umwelt ist damit kein unabhängiger Zufallsschalter je Sensor. Mehrere Verbraucher lesen denselben zeitabhängigen Zustand, setzen ihn jedoch in unterschiedliche Näherungen um. Das Strömungsfeld enthält einen seedbasierten Anteil und windgetriebene Drift. Einige Ozean- und Atmosphärenwerte sind räumlich vereinfacht. Aus einem Windwechsel folgt deshalb keine voll aufgelöste dreidimensionale Meeresströmungsrechnung.')

add('WELT & UMWELT', 'Der Ozean verändert die Schallwege', 'Temperatur, Salzgehalt, Druck und Grenzflächen bilden das akustische Milieu',
    [('Zeitabhängigkeit', 'M2-/S2-Gezeit, Winddurchmischung, Tagesgang, Saison und interne Wellen.'),
     ('Schallgeschwindigkeit', 'Mackenzie-Näherung aus Temperatur, Salzgehalt und Tiefe; Profil bis zum Boden.'),
     ('Grenzflächen', 'Sediment bestimmt Bodenverluste; Wind beeinflusst Oberflächenverluste; Wracks erzeugen Echos.')],
    'Die Thermokline ist kein universeller An/Aus-Schalter: Der gesamte Schallweg entscheidet.',
    ['src/world/ocean.py', 'src/world/world.py', 'src/sonar/propagation.py', 'src/sonar/raytrace.py'],
    'Die Tiefe der durchmischten Deckschicht kombiniert ein räumliches Basisfeld mit Tagesgang, Windmischung und internen Wellen. Das Temperaturprofil und die druckabhängige Schallgeschwindigkeit können ein inneres Minimum erzeugen. Die Strahlverfolgung verwendet eine quantisierte, näherungsweise entfernungsunabhängige Wassersäule. Es handelt sich nicht um einen vollständigen 3-D-Ozeansolver. Gezeiten laufen auf der Ozeanuhr und verändern unter anderem die verfügbare Wassertiefe. Sedimente und Hindernisse werden deterministisch erzeugt.')

add('WELT & UMWELT', 'Wetteranzeige und Flugwetter sind gekoppelt', 'Neue Funktionen des vorliegenden lokalen Arbeitsstands',
    [('Atmosphäre', 'Luftdruck und Tendenz, Temperatur, Böen, Bewölkung, Wolkenuntergrenze, Vereisung, Sonne/Mond.'),
     ('Entscheidungswirkung', 'Helikopter prüft Wind, Querwind, Böen, Sicht, Seegang, Untergrenze und Vereisung.'),
     ('Messgrenze', 'Das Unterwasserprofil im Wetterpanel erscheint erst nach einer BT-Messung und kann veralten.')],
    'Nicht jede atmosphärische Anzeige ist ein eigenes dynamisches Physiksystem; viele Größen sind abgeleitet.',
    ['src/world/atmosphere.py', 'src/ui/weather_station.py', 'src/core/game.py', 'tests/test_weather_station.py'],
    'Game.atmosphere leitet Wetterinstrumente und Himmelsgrößen aus dem bestehenden Umweltzustand ab. Game.helicopter_weather bündelt die operativen Flugwetterentscheidungen. Vereisung kann das Dippen sperren und den Treibstoffbedarf erhöhen. Für Start und Bergung zählt zusätzlich das aktuelle Deckbewegungsfenster. Die Wetter-/Sonaranalyse zeigt das gemessene BT-Profil mit Messalter und Ortsbezug. Die angezeigten Strahlen sind eine Illustration aus dieser Messung und keine Offenlegung des aktuellen wahren Ozeanprofils. Diese Dateien gehören zum lokalen, teils noch nicht eingecheckten Arbeitsstand.')

add('SCHIFF & BEWEGUNG', 'Das Eigenschiff wird kraftbasiert bewegt', 'Befehle verändern Zielwerte; der Rumpf reagiert mit Trägheit',
    [('Längsbewegung', 'Schub minus Widerstand → Beschleunigung. Seegang und zusätzliche Masse verändern die Antwort.'),
     ('Drehbewegung', 'Autopilot → Ruderausschlag → verzögerte Gierrate nach Nomoto → neuer Kurs.'),
     ('Lage & Tiefgang', 'Wellen, Kurvenfahrt und Stabilisierung → Rollen/Stampfen; Masse und Squat → Tiefgang.')],
    'M · dv/dt = a₀ − a₁v − a₂v²; die Längsbewegung wird mit einer geschlossenen Riccati-Lösung integriert.',
    ['src/ship/ship.py', 'src/physics/ship_dynamics.py', 'data/loadouts/ownship_hull.json'],
    'Die Bewegung erfolgt nicht durch sofortiges Setzen der Sollgeschwindigkeit. Die Propeller- und Widerstandsparameter werden an Referenzwerte wie Höchstfahrt, Beschleunigung und Stoppstrecke kalibriert. Das Ruder bewegt sich mit einer begrenzten Stellrate. Die stationäre Gierrate hängt von Fahrt und Ruderausschlag ab; ein stillstehendes Schiff dreht nicht beliebig auf der Stelle. Flutwasser erhöht die Masse. Nach der eigenen Bewegung wird Strömungsdrift hinzugefügt und der Weg auf Grundberührung geprüft. Roll- und Stampfzustände beeinflussen unter anderem den Helikopterbetrieb.')

add('SCHIFF & BEWEGUNG', 'Fahrt ist eine Entscheidung mit vier Folgen', 'Maschinenraum und Brücke verändern die taktischen Möglichkeiten',
    [('Mobilität', 'Höhere Fahrt verkürzt Verlegung, verändert Kurvenfahrt und erhöht die Energie beim Auflaufen.'),
     ('Akustik', 'Fahrt und Kavitation erhöhen Eigenlärm und verschlechtern das passive Hören.'),
     ('Ressourcen & Mittel', 'Treibstoffbedarf, Schleppantenne, Drahtführung und Deckbewegung begrenzen die Nutzung.')],
    'Schneller fahren kann einen Raumgewinn bringen und gleichzeitig den Informationsvorsprung verringern.',
    ['src/ship/ship.py', 'src/sonar/sonar.py', 'src/weapons/torpedo.py', 'src/core/game.py'],
    'Das Schiff koppelt Maschinenzustand und Eigengeräusch über seine aktuelle Fahrt. Ein Leisemodus reduziert den modellierten Lärm. Kavitation verursacht zusätzliche Nachteile. Auch die Schleppantenne besitzt Betriebs- und Fehlerzustände statt einer dauerhaft gültigen Reichweitenverbesserung. Drahtgeführte Torpedos prüfen die Bewegung der Fregatte auf Zugbelastung. Schaden begrenzt erreichbare Fahrt und Steuerfähigkeit. Diese Wechselwirkungen erklären, warum Brücke, Maschinenraum, Sonar und Waffenstation denselben Geschwindigkeitsbefehl unterschiedlich bewerten.')

add('SCHIFF & BEWEGUNG', 'Grundberührung wird entlang des Weges geprüft', 'Auch zwischen zwei Positionen darf der Rumpf keine Küste überspringen',
    [('Geometrie', 'Rumpflänge, Breite, Kurs und dynamischer Tiefgang definieren die Kollisionsform.'),
     ('Bewegungsprüfung', 'Swept Grounding untersucht Translation und Drehung bis zur ersten sicheren Grenzpose.'),
     ('Wirkung', 'Normale Aufprallgeschwindigkeit → kinetische Energie → lokalisierter Schaden und Stillstand.')],
    'E = ½ · m · v²: Eine Verdoppelung der normalen Aufprallfahrt vervierfacht die Aufprallenergie.',
    ['src/world/grounding.py', 'src/ship/ship.py', 'src/ship/damage.py', 'tests/test_grounding_r18.py'],
    'Die Prüfung berücksichtigt den überstrichenen Rumpfbereich und nicht nur die neue Position. Bei Kontakt bleibt das Schiff an einer sicheren Pose stehen; der Kontakt wird eingerastet, damit kein wiederholter identischer Aufprall pro Frame entsteht. Rückwärtsfahrt kann das Schiff lösen, sofern der anschließende Weg frei ist. Die Aufprallenergie wird auf die Komponente senkrecht zur Kontaktfläche bezogen. Tiefgang wächst durch Verdrängung und im Flachwasser durch Squat. Damit hängen Schäden und Gezeiten unmittelbar mit Navigationssicherheit zusammen.')

add('AKUSTIK & SONAR', 'Passives Sonar ist eine Signalkette', 'Ein lautes Ziel ist nicht unter allen Bedingungen leicht hörbar',
    [('Quelle', 'Katalogsignatur + aktuelle Fahrt + Maschinenzustand + transiente Geräusche.'),
     ('Übertragung', 'Entfernung, Absorption, Schallprofil, Boden/Oberfläche und Abschattung.'),
     ('Empfänger', 'Richtwirkung, Eigenlärm, Seegang, Regen und Schiffsverkehr.'),
     ('Kontakt', 'Signalüberschuss erzeugt eine verrauschte Peilung mit Qualität und Zeitstempel.')],
    'Passive Detektion liefert zunächst Richtung und akustische Evidenz; Entfernung muss zusätzlich gewonnen werden.',
    ['src/sonar/sonar.py', 'src/sonar/equation.py', 'src/sonar/propagation.py'],
    'Das Sonarsystem sammelt hörbare Kandidaten einschließlich U-Booten, Schiffen, Tieren, Täuschkörpern und Feindtorpedos. Für jede Quelle bestimmen Signatur, Geometrie und Umwelt die verfügbare Signalenergie. Bei ausreichendem Signal entsteht eine Messung mit Fehler und Qualität. Die interne Sensorerzeugung kennt Kandidaten; diese Kenntnis ist von der dem Spieler zugänglichen Kontaktinformation zu unterscheiden. Die Quellklasse muss aus Messung und Bedienbewertung erschlossen werden. Ein Kontakt ohne gültigen Fix soll keine wahre Zielposition auf der Karte erhalten.', 'flow')

add('AKUSTIK & SONAR', 'Die passive Sonargleichung', 'Signalüberschuss als Bilanz in Dezibel',
    [('SE = SL − TL − NL + DI − DT', 'SL: Quellpegel\nTL: Übertragungsverlust\nNL: Rauschpegel\nDI: Richtgewinn\nDT: Detektionsschwelle'),
     ('Was verändert die Bilanz?', 'Fahrt und Kavitation erhöhen Eigenlärm. Frequenz und Weg bestimmen Verlust. Antenne und Empfindlichkeit verändern den nutzbaren Pegel.')],
    'Die Modelle werden an Spielreferenzen kalibriert; eine nominelle Reichweite ist kein fester Detektionskreis.',
    ['src/sonar/equation.py', 'src/sonar/sonar.py', 'tools/calibrate.py'],
    'Die kanonischen Ausbreitungsbänder liegen bei 100, 400, 1600 und 6400 Hz. Das Rauschen kombiniert Wind/Seegang, entfernte Handelsschifffahrt, Regen sowie Eigen- und thermisches Rauschen. Dezibelwerte werden für unabhängige Rauschleistungen in lineare Leistungen übersetzt, summiert und wieder logarithmiert. Absorption folgt einer Francois-Garrison-Näherung. Die Referenzkonstanten sind an bestehende Spielreichweiten angepasst. Der praktische Vergleich lautet deshalb: Welche Größe ändert den Signalüberschuss in dieser Situation?', 'cards')

add('AKUSTIK & SONAR', 'Strahlverfolgung statt pauschaler Schichtstrafe', 'Die Form des Schallprofils erzeugt unterschiedliche Empfangszonen',
    [('Eingabe', 'Quantisiertes Profil, Sendertiefe, Wassertiefe, Sediment und Wind.'),
     ('Berechnung', '48 Strahlen; Brechung nach Snell; Umkehrpunkte; Oberflächen- und Bodenreflexion.'),
     ('Ergebnis', 'Verlusttabellen über Entfernung/Tiefe je Band; begrenzter Cache mit 64 Einträgen.')],
    'Schatten, Oberflächenkanal und Mehrwegeausbreitung entstehen im vereinfachten Profilmodell.',
    ['src/sonar/raytrace.py', 'src/sonar/propagation.py', 'src/world/ocean.py'],
    'Das Modell verfolgt einen begrenzten Strahlenfächer und summiert Energie inkohärent in einem Entfernungs-Tiefen-Raster. Oberflächenrauigkeit und Sediment dämpfen Reflexionen. Es berechnet keine vollständige Wellengleichung und keinen kohärenten Interferenzraum. Der Umweltschlüssel wird quantisiert, damit ähnliche Situationen dieselbe Tabelle verwenden. Der Cache muss eine reine Optimierung bleiben: Ein Treffer oder Neuaufbau darf das physikalische Ergebnis nicht von der Aufrufreihenfolge abhängig machen. Die Tabellen reichen bis zur implementierten Modellgrenze von 80 NM.')

add('AKUSTIK & SONAR', 'Bug-, Schlepp- und abgesetzte Sensoren', 'Die Position des Empfängers ist Teil der Messung',
    [('Bugsonar', 'Direkt am Eigenschiff; Fahrt, Eigenlärm, Rumpf und Schäden beeinflussen den Empfang.'),
     ('Schleppantenne', 'Eigener Ausbringzustand und Tiefe; Richtwirkung und Links/Rechts-Mehrdeutigkeit.'),
     ('Dippen & Bojen', 'Andere Beobachterpositionen liefern zusätzliche Peilungen und verbessern die Geometrie.')],
    'Mehr Sensoren helfen vor allem dann, wenn sie neue, verwertbare Beobachtungsgeometrie liefern.',
    ['src/sonar/sonar.py', 'src/air/helicopter.py', 'src/air/sonobuoy.py', 'src/audio/receiver.py'],
    'Eine lineare Schleppantenne kann spiegelbildliche Peilungen erzeugen. Ungeklärte Seitenmehrdeutigkeit wird nicht einfach als sichere TMA-Evidenz behandelt; eigene Manöver oder das Bugsonar können helfen. Das Antennenmodell enthält eine sinc-artige Haupt- und Nebenkeulencharakteristik. Abgesetzte Sensoren haben eigene Positionen und Tiefen. Sie sind nicht mit demselben Rumpfeigenlärm wie das Bugsonar belastet. Bojen und Dippsonar speisen ihre Beobachterpositionen in die Peilungshistorie ein. Das ist gemeinsame Schätzung aus mehreren Empfängern, kein allgemeiner bistatischer Aktivsonarsimulator.')

add('AKUSTIK & SONAR', 'Aktivsonar tauscht Tarnung gegen einen Fix', 'Ein Ping verursacht ein zeitverzögertes Messereignis',
    [('Aussenden', 'Pulswahl CW oder LFM, Bereitschaft und Sensorzustand werden berücksichtigt.'),
     ('Weg & Echo', 'Hin- und Rückweg, Zielstärke, Nachhall, Doppler und Schallgeschwindigkeit bestimmen Empfang.'),
     ('Publizieren', 'Ein gespeichertes Echo trifft später ein und liefert fehlerbehaftete Peilung, Entfernung und Tiefe.'),
     ('Gegenwirkung', 'Andere Einheiten können den Ping bemerken; Wracks können unzugeordnete Echos liefern.')],
    'Echolaufzeit ≈ 2R/c. Bei 10 NM und c = 1500 m/s sind das rund 24,7 Sekunden.',
    ['src/sonar/sonar.py', 'src/sonar/equation.py', 'src/world/world.py', 'tests/test_active_sonar_snapshot.py'],
    'Das Zahlenbeispiel ist eine Umrechnung mit nominaler Schallgeschwindigkeit, keine gemessene Szene. Das aktive Signalbudget enthält den doppelten Übertragungsverlust und die aspektabhängige Zielstärke. CW kann bewegte Ziele über Doppler besser von Nachhall trennen; LFM hat Bandbreiten- und Verarbeitungsgewinn. Ausgesendete Pings und ausstehende Echos tragen Schnappschussdaten, damit ein späteres Echo keine zukünftige Position verrät. Deren Zustand muss auch über Speichern/Laden erhalten bleiben.', 'flow')

add('AKUSTIK & SONAR', 'TMA gewinnt Bewegung aus einer Peilungsfolge', 'Target Motion Analysis schätzt; sie liest keine Zielkoordinaten',
    [('Messhistorie', 'Zeit, rohe Peilung, Messfehler und eigene Beobachterposition; optional Doppler.'),
     ('Schätzer', 'Grobe Kurs-/Fahrtsuche, lokale Verfeinerung und Levenberg-Marquardt-Anpassung.'),
     ('Lösung', 'Geschätzte Position, Kurs, Fahrt, Qualität und Kovarianzellipse; Hysterese stabilisiert Neulösungen.')],
    'Ohne ausreichende Geometrie kann es viele plausible Lösungen geben – auch bei einer langen Kontaktspur.',
    ['src/sonar/tma.py', 'src/sonar/tma_lm.py', 'src/sonar/sonar.py', 'tests/test_arrays_doppler_tma.py'],
    'Der Schätzer verlangt Mindestdatenmenge und Zeitspanne sowie Beobachtbarkeit. Eine Kursänderung ohne räumliche Bewegung löst die Entfernungsmehrdeutigkeit nicht automatisch. Reale Abweichung der Eigenbewegung von einer geraden gleichförmigen Bahn ist entscheidend; Doppler kann zusätzliche Information liefern. Der Schätzer verwendet die Rohmessungen, während Anzeigeglättung getrennt bleibt. Hysterese hält bei fast gleichwertigen Kandidaten die vorherige Lösung. Die Kovarianz beschreibt Unsicherheit innerhalb des vereinfachten Bewegungsmodells und ist keine Garantie auf wahre Zielnähe.')

add('AKUSTIK & SONAR', 'Was die Sonaranzeigen tatsächlich zeigen', 'Verschiedene Auswertungen derselben beobachteten Signalwelt',
    [('Breitband', 'Energie über Peilung und Zeit', 'Wo treten akustische Quellen auf?'),
     ('LOFAR', 'Schmalbandlinien über Frequenz/Zeit', 'Welche stabilen tonalen Merkmale sind hörbar?'),
     ('DEMON', 'Modulation aus der Signalhüllkurve', 'Welche Periodizitäten und RPM-Hypothesen passen?'),
     ('TMA', 'Geometrie und geschätzte Bewegung', 'Wo könnte der Kontakt mit welcher Unsicherheit sein?'),
     ('BT / Umwelt', 'Gemessenes Temperatur-/Schallprofil', 'Welche Empfangsbedingungen sind plausibel?')],
    'Signaturähnlichkeit ist Evidenz; Klassifikation und Zugehörigkeit bleiben eine Bedienbewertung.',
    ['src/audio/receiver.py', 'src/audio/demon.py', 'src/sonar/sonar.py', 'src/ui/sonar_view.py'],
    'LOFAR und DEMON werden aus synthetisierten Empfängersignalen abgeleitet. DEMON arbeitet auf deren Hüllkurve und liefert Hypothesen, etwa für Blattzahl, Harmonische und Drehzahl. Es liest keine verborgene wahre Propelleridentität aus der Plattform. Die Signaturen sind illustrative Modellparameter, keine realen Tonaufnahmen. BT ist eine eigene Messung mit Fehler und Gültigkeitsbezug. Der Kontaktanalysator stellt Katalogreferenzen zum Vergleichen bereit; dadurch wird eine aktuell gehörte Quelle nicht automatisch eindeutig identifiziert.', 'table')

add('AKUSTIK & SONAR', 'Sonar als Arbeitsplatz', 'Repository-Abbildung: mehrere Auswertungen unterstützen eine gemeinsame Entscheidung',
    [('docs/screenshots/de-sonar-tma.png', 'TMA verbindet Peilungshistorie, Eigenbewegung und Unsicherheit.')],
    'Die Anzeige ist eine Sicht auf Messungen und Schätzungen; der wahre Zielzustand bleibt im Simulationskern.',
    ['src/ui/sonar_view.py', 'src/sonar/tma.py', 'docs/screenshots/de-sonar-tma.png'],
    'An dieser Abbildung lässt sich die Rolle des Bedieners erklären: Kontakt auswählen, Historie bewerten, eigene Manöver planen und die Lösung vor einer Waffenfreigabe überprüfen. Eine gut aussehende Kurve allein macht noch keinen frischen und belastbaren Schussdatensatz. Der Screenshot stammt aus dem Repository und illustriert den Arbeitsplatz; er ist kein neu aufgezeichneter Lauf des aktuellen Arbeitsstands.', 'screen')

add('SENSOREN & LAGE', 'Radar: Reichweite, Suchstrahl und Wahrscheinlichkeit', 'Ein Kontakt erhält erst beim Strahldurchgang einen neuen Look',
    [('Signalmodell', 'Echo ∝ Radarquerfläche / R⁴. Seegang, Regen und Störung verringern das Signal-Stör-Verhältnis.'),
     ('Detektionsmodell', 'Swerling-1-Fluktuation und CA-CFAR-Näherung liefern eine Detektionswahrscheinlichkeit.'),
     ('Zeit & Geometrie', 'Antenne dreht mit 90°/s; Abschattung und Radarhorizont begrenzen relevante Ziele.')],
    'Ein deterministischer Zufallsentscheid pro Look macht Detektion reproduzierbar, aber nicht garantiert.',
    ['src/sensors/radar.py', 'src/core/game.py', 'src/core/detrand.py', 'tests/test_radar_emissions_ecm.py'],
    'Bei 90 Grad pro Sekunde braucht eine vollständige Umdrehung vier Simulationssekunden. Die Publikationskadenz ist also nicht gleich Wiederholrate eines bestimmten Ziels. Für ein Referenzziel ist die nominelle Reichweite auf eine Entdeckungswahrscheinlichkeit von 0,5 kalibriert. Ein Rauschstörer fällt auf dem Hinweg ungefähr mit R^-2 ab, das Echo auf dem Hin- und Rückweg mit R^-4; dadurch entsteht ein modellierter Burn-through-Bereich. Zielaspekt verändert die Radarrückstrahlung. Abwehr und Lagebild verwenden die zuletzt tatsächlich erzeugten Messungen.')

add('SENSOREN & LAGE', 'ESM beobachtet Emissionen, ECM verändert Suchköpfe', 'Passive Aufklärung und aktive Störung haben unterschiedliche Rollen',
    [('ESM / ELOKA', 'Misst Peilung, Frequenz, Puls-/Scanmerkmale und Pegel; vergleicht Signale mit Emittentenkandidaten.'),
     ('Betriebszustand', 'Rotierende Sender erzeugen zeitliche Auffassungen; Tracks können live, gehalten oder veraltet sein.'),
     ('ECM', 'Frische Signale und begrenzte Kanäle steuern modellierte Rausch- und Täuschverfahren gegen Suchköpfe.')],
    'Eine ESM-Reichweitenschätzung hängt von Annahmen über die Senderleistung ab und ist kein Radarfix.',
    ['src/sensors/esm.py', 'src/core/game.py', 'src/air/asm.py'],
    'ESM empfängt ausgesandte Signale statt eigener Echos. Ein passendes Frequenz- oder Pulsmuster liefert Kandidaten mit Güte, keine allwissende Identifikation. Pegel und angenommene Leistungsklasse erlauben eine grobe Entfernungsabschätzung. Die eigenen ECM-Techniken greifen in die terminale Suchkopfentscheidung ein. Feindliche Flugkörper können einen Rauschstörer tragen. Ein aktiver ASM-Suchkopf selbst ist eine ESM-Emission und kann damit eine Warnung auslösen, bevor das Suchradar einen stabilen Körpertrack hat. Die genaue Warnreihenfolge bleibt geometrie- und sensorabhängig.')

add('SENSOREN & LAGE', 'AIS, Ausguck und HFDF ergänzen sich', 'Jeder Sensor beantwortet andere Fragen und hat eigene Grenzen',
    [('AIS', 'Empfangene dynamische und statische Meldungen', 'VHF-Sichtlinie, Meldeintervall, Gelände'),
     ('Ausguck', 'Sichtkontakt und begrenzte beobachtete Merkmale', 'Kontrast, Zielgröße, Horizont, Sicht und Mondlicht'),
     ('HFDF', 'Peilung eines Funkereignisses plus Frequenz', 'Boden-/Raumwelle, Sprungzone, Tag/Nacht'),
     ('Funk / HQ', 'Meldungen und grobe Aufklärungshinweise', 'Ein Hinweis ist kein frischer Sensorfix')],
    'Ein AIS-Name, eine Funkpeilung und ein Radarpunkt sind unterschiedliche Informationsprodukte.',
    ['src/sensors/ais.py', 'src/sensors/visual.py', 'src/sensors/hfdf.py', 'src/core/game.py'],
    'Der simulierte AIS-Empfänger erhält Klasse-A-artige Positionsmeldungen mit fahrtabhängigen Intervallen und statische Angaben in längeren Abständen. Radar allein darf nicht den wahren Schiffsnamen oder Kurs liefern. Der Ausguck verwendet eine Kontrast- und Winkelgrößennäherung mit optischem Horizont. HFDF berücksichtigt eine vereinfachte ionosphärische Ausbreitung, optimale Arbeitsfrequenz und Sprungdistanz; Raumwellenpeilungen haben größere Fehler. Der anfängliche HQ-Hinweis ist bewusst gerundet und unbestätigt. Zusammengenommen ergänzen sich diese Quellen, ohne dieselbe Genauigkeit zu besitzen.', 'table')

add('SENSOREN & LAGE', 'Ein Track altert – auch wenn sein Symbol stehen bleibt', 'Messzeit, Positionszeit und Darstellungsfilter sind getrennt',
    [('Rohmessung', 'Peilung, Entfernung und Messzeit werden als begrenzte Historie aufbewahrt.'),
     ('Darstellung', 'Glättung und Bewegungsfit beruhigen die Anzeige; sie erzeugen keine neue Messung.'),
     ('Gültigkeit', 'Quelle, Alter, Qualität und Positionsfrische entscheiden über Anzeige und Waffenverwendung.')],
    'Eine geglättete Position darf eine alte Waffenlösung nicht künstlich verjüngen.',
    ['src/sensors/tracks.py', 'src/ui/observations.py', 'src/sonar/sonar.py', 'tests/test_observation_smoothing.py'],
    'SensorTrack trennt unter anderem raw_x/raw_y, dargestellte Position, last_seen und position_seen. Neue Richtungsinformation kann vorhanden sein, obwohl kein neuer Entfernungsfix vorliegt. Bewegungsrichtung und Fahrt können aus einer Ausgleichsgeraden durch mehrere Positionsmessungen abgeleitet werden. Diese Schätzung ist wiederum von einem direkt gemeldeten AIS-Kurs zu unterscheiden. Anzeigeglättung arbeitet auf Präsentationsdaten; TMA und Feuerleitung sollen auf geeignete Rohdaten und echte Frischekriterien zurückgreifen. Die einzelnen Sensorprodukte besitzen unterschiedliche Verfallsregeln.')

add('SENSOREN & LAGE', 'OPZ-Fusion ist eine bewusste Zuordnung', 'Der Bediener verbindet Beobachtungen; die Quellen bleiben erhalten',
    [('Auswählen', 'Beobachtungen markieren und als zusammengehörig bewerten.'),
     ('Zusammenfassen', 'Vorhandene Positionen werden nach Qualität gewichtet; Peilungen zirkulär gemittelt.'),
     ('Annotieren', 'Klasse und Zugehörigkeit setzen; falsche Verknüpfungen wieder lösen.')],
    'Nur Peilungen zu fusionieren erzeugt hier nicht automatisch eine triangulierte Entfernung.',
    ['src/sensors/fusion.py', 'src/core/game.py', 'tests/test_opz_fusion.py'],
    'OPZFusionPicture ist ein begrenzter, transienter Arbeitsbereich. Eine Fusion referenziert Beobachtungs-IDs und überschreibt die Quellsysteme nicht. Ihre Position entsteht nur aus Mitgliedern, die bereits positioniert sind. Sind alle Quellen reine Peilungen, bleibt die Position unbekannt; es wird keine allgemeine Kreuzpeilungsrechnung durchgeführt. Die Fusion übernimmt auch nicht automatisch den wahren Kurs oder die Klasse eines Ziels. Manuelle Klassifikation und Zugehörigkeit sind Bedienerannotation. Deshalb muss die spätere Zielverwendung weiterhin ihre eigenen Gültigkeitsprüfungen durchlaufen.')

add('GEGNER & VERKEHR', 'Auch NPCs haben ein eigenes Lagebild', 'Wahrnehmen, erinnern, entscheiden, handeln',
    [('Sensoren', 'PlatformSensorSuite prüft profilierte Sensoren, Betriebszustände, Schäden und Gegenkandidaten.'),
     ('Gedächtnis', 'Abgetrennte Beobachtungen und alternde Erinnerungen bestimmen die taktische Grundlage.'),
     ('Verhalten', 'Zustandslogik wählt Patrouille, Ausweichen, Lauern, Angriff oder Rückkehr-/Auftauchphasen.')],
    'Der normale NPC-Pfad nutzt Beobachtungen; einzelne Legacy-/Kompatibilitätspfade bleiben gesondert vorhanden.',
    ['src/sensors/platform.py', 'src/enemies/sub.py', 'src/enemies/surface.py', 'src/core/game.py'],
    'Game aktualisiert die NPC-Sensoren in stabiler Reihenfolge und vor deren Entscheidungen. Ein U-Boot muss aus passiven Peilungen zunächst eigene Entfernungsinformation gewinnen. Es kann Ping- und Torpedoereignisse bemerken, Kontaktgedächtnis führen und nach einer Reaktionszeit ausweichen. Feindliche Gruppen können modellierte Datalink-Beobachtungen austauschen, bei getauchten Booten jedoch nur unter Erreichbarkeitsbedingungen. Der Code enthält weiterhin Legacy- bzw. direkte Testaufrufpfade mit vereinfachten Beobachtungssnapshots; die Präsentation behauptet daher keine absolute Gleichheit sämtlicher Pfade.')

add('GEGNER & VERKEHR', 'U-Boot-Energie erzwingt taktische Phasen', 'Batterie, AIP und Schnorcheln beeinflussen Ortbarkeit und Bewegung',
    [('Getaucht', 'Bordlast + fahrtabhängiger Antriebsbedarf entnehmen Energie; verfügbare Leistung begrenzt Fahrt.'),
     ('Reserve erreicht', 'AIP unterstützt, sofern vorhanden; sonst Aufstieg zur Schnorcheltiefe.'),
     ('Exponierte Phase', 'Schnorcheln lädt; eine Funkphase kann folgen; anschließend Rückkehr zur Tauchtiefe.')],
    'Die Energielage erzeugt Situationen für Radar, ESM, HFDF und akustische Aufklärung.',
    ['src/enemies/endurance.py', 'src/enemies/sub.py', 'src/physics/submarine.py'],
    'SubmarineEndurance ist ein fiktives, aber explizites Energiemodell mit Last, Generator, Batterie und gegebenenfalls endlichem AIP-Vorrat. Reservegrenzen steuern die Phasen SUBMERGED, AIP, ASCENDING, SNORKEL, RADIO und DESCENDING. Nicht jedes Profil verwendet dieselbe Ausstattung. Die Bewegung ergänzt beschränkte Beschleunigung, fahrtabhängige Tiefenruderwirkung, druckabhängige Kavitation und Rumpfbelastung. Schaden kann Lärm erhöhen und Notauftauchen erzwingen. Taktische und energetische Zustandsautomaten greifen dadurch ineinander.')

add('GEGNER & VERKEHR', 'Oberfläche, Tiere und Luftverkehr formen die Lage', 'Nicht jedes hörbare oder sichtbare Objekt ist ein Angriffsziel',
    [('Überwasserschiffe', 'Zivile Routen und Kampfschiffverhalten; Bewegung, Emissionen, Sensoren und gegebenenfalls Waffen.'),
     ('Biologische Quellen', 'Eigene Bewegung und akustische Signaturen erzeugen Kontakte ohne militärischen Zielwert.'),
     ('Luftverkehr', 'Transitflüge, gesonderte Angreifer und optional Live-Flugzeuge besitzen unterschiedliche Modelle.')],
    'Dichte Lage erhöht Klassifikationsaufwand, akustischen Hintergrund und das Risiko einer Fehlbekämpfung.',
    ['src/enemies/surface.py', 'src/enemies/animal.py', 'src/air/flights.py', 'src/air/raid.py'],
    'Oberflächenkampfschiffe reagieren auf ihre sensorisch gewonnene Lage und können Flugkörper oder ASW-Waffen einsetzen, sofern Profil und Bereitschaft dies erlauben. Zivile Schifffahrt trägt sowohl zum sichtbaren Lagebild als auch zum modellierten Hintergrundgeräusch bei. Tiere besitzen eigene Bewegung und Signaturen, wodurch die Sonaraufgabe mehr als eine reine Gegnerliste wird. Reguläre Transitflüge verwenden vereinfachte Flugbahnen; Raider besitzen ein Angriffsprofil mit begrenztem Kurvenflug und Pop-up-Phase. Diese Klassen sollten nicht als einheitlicher Flugsimulator beschrieben werden.')

add('WAFFEN & WIRKUNG', 'Waffenfreigabe ist eine Kette von Bedingungen', 'Ein ausgewähltes Symbol allein genügt nicht',
    [('Lage', 'Kontakt/Track vorhanden, passende Domäne, ausreichend frische Beobachtung oder Lösung.'),
     ('Einsatzregel', 'Bedienentscheidung und ROE; bei Remote Crew zusätzlich Stationsrecht und direkte Feuerfreigabe.'),
     ('Bereitschaft', 'Waffenstation, Munition, Rohr-/Magazinzustand, Ladezeit und Plattformbereitschaft.'),
     ('Geometrie', 'Reichweite, Tiefe, Startpunkt und jeweilige Waffen- bzw. Abwurfhülle.')],
    'Die konkreten Bedingungen sind waffenspezifisch und werden beim tatsächlichen Befehl erneut geprüft.',
    ['src/core/game.py', 'src/weapons/asw.py', 'src/weapons/air_defense.py', 'src/commander/bridge.py'],
    'Die Präsentation fasst die gemeinsamen Kategorien zusammen, ohne eine einzige universelle boolesche Prüfung zu behaupten. Torpedos benötigen andere Beobachtungsdaten als ESSM oder CIWS. Eine verlorene Sensorposition, eine beschädigte Station oder ein leeres Rohr kann einen zuvor plausiblen Angriff verhindern. Das Waffensystem modelliert Lade- und Verbrauchszustände separat von der Zielauswahl. Remote-Crew-Befehle müssen kurz vor ihrer Ausführung nochmals gegen den aktuellen Zustand geprüft werden, weil sich zwischen Browseranzeige und Empfang die Lage geändert haben kann.', 'flow')

add('WAFFEN & WIRKUNG', 'Torpedo: Start, Draht, Suche und Endphase', 'Die Führung wechselt ihren Informationszugang im Flugverlauf',
    [('Start', 'Rohrbereitschaft, Motorhochlauf, Starttiefe und hörbarer Ausstoßtransient.'),
     ('Mittellauf', 'Fahrt zum befohlenen Datum; frische Kontaktfixes können über den Draht aktualisieren.'),
     ('Terminalphase', 'Suchkopf prüft erreichbare Kandidaten und kann Ziel oder Täuschkörper auffassen.'),
     ('Wirkung', 'Nächste Annäherung und Tiefengeometrie → Zündung/Schaden; Boden und Energie begrenzen den Lauf.')],
    'Vor der Eigenortung folgt die Waffe einem befohlenen Datum; danach wirkt ihr lokaler Suchkopf.',
    ['src/weapons/torpedo.py', 'src/physics/torpedo_dyn.py', 'src/core/game.py'],
    'Der reguläre Spielpfad gibt Torpedos eine beobachtete Zielposition mit. Drahtupdates beruhen auf einer frischen Kontaktlösung, nicht fortlaufend auf der wahren Position des ursprünglich gewählten Gegners. Nähe zum Datum aktiviert die terminale Suche. Dort darf der Suchkopf intern reale Kandidaten auf geometrische und akustische Eignung untersuchen. Zielwahl ist damit nicht zwangsläufig dauerhaft an den Startkontakt gebunden. Der Torpedocode besitzt auch Fallbacks für Aufrufe ohne Datum; diese sind von der regulären beobachtungsbasierten Einsatzkette zu unterscheiden.', 'flow')

add('WAFFEN & WIRKUNG', 'Warum ein Torpedo verfehlen kann', 'Bewegung, Energie, Verbindung und Zielwahl setzen eigene Grenzen',
    [('Dynamik & Energie', 'Motorhochlauf, fahrtabhängige Dreh- und Tiefenreaktion; Leistungsbedarf ungefähr ∝ v³.'),
     ('Draht', 'Zwei endliche Spulen; zu hohe Fahrt oder anhaltend harte Kurven können die Verbindung brechen.'),
     ('Täuschung & Wirkung', 'Suchkopf konkurriert um Signale; Annäherungsschaden hängt vom Abstand ab statt nur von HIT/NO HIT.')],
    'Aus einer guten Startlösung folgt keine garantierte Zielvernichtung.',
    ['src/weapons/torpedo.py', 'src/physics/torpedo_dyn.py', 'src/enemies/decoy.py', 'src/weapons/asw.py'],
    'Bei leerem Energiespeicher endet der Motorlauf; die Waffe kann zunächst auslaufen und wird erst bei zu geringer Fahrt verloren. Eigene und gegnerische Torpedos haben Tiefenregelung und geschwindigkeitsabhängige Kurvenbegrenzung. Drahtlänge und Zugbelastung hängen auch von der Bewegung der Fregatte ab. Täuschkörper werden nicht nur als pauschale Ausweichchance betrachtet: Ihre Signale nehmen an der terminalen Zielwahl teil. Die Schadensnähe folgt einer Stoßfaktor-Näherung proportional zu sqrt(Sprengstoffmasse)/Abstand. Die genaue Wirkung ist zielklassenspezifisch.')

add('WAFFEN & WIRKUNG', 'ASROC und Helikopter verändern den Zustellweg', 'Die Suchwaffe bleibt an ein beobachtetes Datum gebunden',
    [('Schiffstorpedo', 'Start am eigenen Schiff → längerer Unterwasseranlauf → drahtgestützte Updates möglich.'),
     ('ASROC', 'Flug zum befohlenen Einwasserpunkt → Torpedopayload → Helixsuche und Tiefenregelung.'),
     ('Helikoptertorpedo', 'Verlegung des eigenen Mittels → gültige Abwurfgeometrie → autonome Unterwassersuche.')],
    'Andere Zustellung erweitert die Einsatzmöglichkeiten, beseitigt aber keine Unsicherheit des Zieldatums.',
    ['src/weapons/asw.py', 'src/air/helicopter.py', 'src/core/game.py'],
    'ASROC besitzt eine eigene Flugphase. Erst bei erfolgreichem Einwassern wird die Unterwasserwaffe erzeugt. Deren Bewegung beginnt im nächsten Physikschritt, sodass Zeit nicht doppelt verbraucht wird. Der Helikopter berechnet eine relative Abwurfgeometrie aus einem beobachteten Datum; er darf nicht direkt zu einer verborgenen aktuellen Zielposition geführt werden. Entfernung, Frische und Bereitschaft werden im jeweiligen Freigabepfad geprüft. So bleibt die Beobachtungslage trotz verschiedener Zustellplattformen die gemeinsame Grundlage.')

add('LUFT & ABWEHR', 'Ein Luftangriff besteht aus mehreren Phasen', 'Angreifer, Flugkörper und Suchkopf sind getrennte Zustände',
    [('Angreifer', 'Begrenzter Kurven-/Steigflug; Pop-up auf Angriffshöhe und Feuerleitphase.'),
     ('Flugkörper', 'Boost, Marschflug und Höhenregelung; zunächst inertial zum Startdatum.'),
     ('Suchkopf', 'Aktivierung, Sichtfeld und Lock-Zeit; anschließend begrenzte Proportionalnavigation.')],
    'Radarhorizont, Emissionen und Flugprofil bestimmen, wann welche Warnung möglich wird.',
    ['src/air/raid.py', 'src/air/asm.py', 'src/physics/missile.py', 'src/core/game.py'],
    'Raider sind keine voll simulierten fliegenden Cockpits, besitzen aber koordinierte Kurven und begrenzte Vertikalbewegung. Die Pop-up- und Feuerleitphase macht sie anders wahrnehmbar als den tief fliegenden Flugkörper. ASM werden als begrenzte Punktmassen mit Schub- und Höhenphasen geführt. Im Endanflug beschränken Suchfeld, Aufschaltzeit und Querbeschleunigung die Verfolgung. Eigene ECM und Düppel können die Suchentscheidung beeinflussen. Die ESM-Emission des aktiven Suchkopfs ist ein zusätzlicher Warnkanal neben der Radarauffassung des Flugkörpers.')

add('LUFT & ABWEHR', 'Luftabwehr ist ein begrenztes Schichtsystem', 'Softkill und Hardkill greifen in denselben Angriff ein',
    [('ECM / Düppel', 'Störung und Scheinziele beeinflussen den Suchkopf; Wolken entfalten sich, driften und zerfallen.'),
     ('ESSM', 'Start gegen einen frischen Track; Vorsteuerung aus Beobachtungen, danach eigener Suchkopf.'),
     ('CIWS / FLAK', 'Freigaben, Munition, Bereitschaft und Nahbereich entscheiden; CIWS muss nachführen und rechtzeitig wirken.')],
    'Mehrere gleichzeitige Angriffe können Feuerkanäle, Nachführung und Munitionsvorräte sättigen.',
    ['src/air/chaff.py', 'src/air/asm.py', 'src/weapons/ciws.py', 'src/weapons/air_defense.py', 'src/core/game.py'],
    'Düppel sind gespeicherte Wolken mit Alter, Bloom und Winddrift. Das Suchkopfmodell vergleicht deren Echo mit dem Schiff und kann auf eine Wolke umlenken. ESSM erhält vor der Eigenortung Beobachtungsupdates. CIWS besitzt eine begrenzte Schwenkrate; die Burst-Wirkung berücksichtigt Streuung, Vorhersagefehler und Flugzeit der Geschosse. Ein Burst, der erst nach dem Einschlag eintreffen würde, kann den Angriff nicht mehr abwehren. Die Abwehr folgt den aktuellen Freigaben und begrenzten Ressourcen, nicht einer globalen Unverwundbarkeitswahrscheinlichkeit.')

add('LUFT & ABWEHR', 'Der Helikopter verbindet drei Stationen', 'Brücke schafft Bedingungen, Sonar gewinnt Daten, Waffen nutzen das Datum',
    [('Flugbetrieb', 'Hangar → Einsatz → Rückkehr; Wegpunkte, Treibstoff, Wetter und Deckfenster.'),
     ('Aufklärung', 'Schwebeflug, Dipptiefe und Bojen erzeugen Beobachtungen von anderen Orten.'),
     ('Ressourcen', 'Schweben benötigt mehr Treibstoff; Vereisung, Vorräte und beschädigtes Flugdeck beschränken Einsätze.')],
    'Die Fregatte kann durch Kurs und Fahrt die Bedingungen für Start, Bergung und Sensorik verändern.',
    ['src/air/helicopter.py', 'src/air/sonobuoy.py', 'src/core/game.py'],
    'Der Helikopter ist ein kommandiertes eigenes Mittel und daher mit seiner eigenen Position sichtbar. Das bedeutet jedoch keinen Zugriff auf verborgene Feindpositionen. Beim Dippen wirkt Wind gegen eine vereinfachte Positionshaltung; der Treibstoffverbrauch ist gegenüber dem Reiseflug erhöht. Sonarbojen driften mit Oberflächenströmung und einem zusätzlichen Windanteil. Das Flugdeck benötigt ein Roll-/Stampffenster, im Modell 8 Grad Rollen und 3,5 Grad Stampfen. Wind, Sicht, Seegang und neue Vereisungsregeln ergänzen diese mechanischen Grenzen.')

add('SCHADEN & RÜCKKOPPLUNG', 'Ein Treffer setzt eine Folge von Zuständen in Gang', 'Neun fiktive Zonen verbinden Schadensort und Systemfunktion',
    [('Initialwirkung', 'Trefferlage wählt betroffene Abteilungen; Torpedo/Grundkontakt erzeugen Lecks, ASM vor allem Brand.'),
     ('Ausbreitung', 'Wasserdruck treibt Flutung; Wärme, Brennstoff und Sauerstoff treiben Feuer und Übergreifen.'),
     ('Systemwirkung', 'Leistungsfähigkeit sinkt; Masse, Tiefgang und Krängung verändern Bewegung und Überlebensfähigkeit.')],
    'Schaden verändert die gesamte Einsatzkette: Sensoren, Waffen, Antrieb, Steuerung und Flugdeck.',
    ['src/ship/damage.py', 'src/core/game.py', 'src/ship/ship.py'],
    'Die neun Zonen sind ein fiktives funktionales Abteilungsmodell und kein realer Bauplan. Ein Torpedotreffer ordnet anhand der Aufpralllage lokale Räume zu. Flutung kann elektrische Kurzschlüsse auslösen; anhaltend heiße Schotten können Feuer weitertragen. Ein Magazinbrand kann eine Sekundärwirkung erzeugen. Stationen besitzen eine kontinuierliche capability zusätzlich zu Zuständen wie zerstört. Im nächsten Navigationsschritt werden daraus Fahrtgrenze, Steuerungszustand, Stabilisatorzustand und Flutwassermasse übernommen. So schließen sich die physikalischen Rückkopplungen.')

add('SCHADEN & RÜCKKOPPLUNG', 'Flutung koppelt Hydraulik und Stabilität', 'Die Lage des Wassers ist ebenso wichtig wie seine Menge',
    [('Wassereinbruch', 'Q = Cᵈ · A · √(2gh)\nLeckfläche und Druckhöhe bestimmen den Zustrom.'),
     ('Hydrostatik', 'Wassermasse erhöht Verdrängung und Tiefgang; freie Oberflächen vermindern Stabilität.'),
     ('Versagen', 'Seitliche Wassermomente erzeugen Schlagseite; verlorene Reserve oder Stabilität führen zum Sinken/Kentern.')],
    'Mehr Tiefgang kann den Wasserdruck erhöhen – eine physische Rückkopplung innerhalb des Schadensmodells.',
    ['src/ship/damage.py', 'src/ship/ship.py', 'tests/test_damage_stability.py'],
    'Die Druckhöhe wird zwischen äußerer Wasserlinie und innerem Wasserstand bestimmt. Deshalb verlangsamt sich ein Einbruch beim Niveauausgleich; hoch liegende Räume werden nicht automatisch wie Räume unter der Wasserlinie geflutet. Jede Zone besitzt Volumen, Schwerpunkt und Höhenlage. Freie Flüssigkeitsoberflächen reduzieren die metazentrische Stabilität. Das Modell prüft Reserveauftrieb, GM und einen Krängungsgrenzwert. Diese Größen sind spielkalibrierte Näherungen und keine vollständige hydrostatische Schiffsentwurfsrechnung.')

add('SCHADEN & RÜCKKOPPLUNG', 'Reparatur ist eine logistische Entscheidung', 'Personal und Material brauchen Zeit am richtigen Ort',
    [('Drei Trupps', 'Zuweisung an Abteilungen; Weg über den Abteilungsgraphen mit gespeicherter Ankunftszeit.'),
     ('Begrenztes Material', 'Lecks werden zuerst mit Patch-Kits verkleinert; Restleckage kann bestehen bleiben.'),
     ('Prioritäten', 'Pumpen, Brandbekämpfung und Wiederherstellung konkurrieren um dieselben Ressourcen.')],
    'Eine Sensorreparatur kann taktisch wertvoll sein; eine Stabilitätskrise kann sie dennoch überstimmen.',
    ['src/ship/damage.py', 'src/core/autocrew.py', 'tests/test_damage.py'],
    'Reparatur beginnt erst nach Ankunft des Trupps. Die Laufzeit beträgt im Modell 20 Sekunden pro Graphkante; acht Patch-Kits stehen anfänglich zur Verfügung. Ein provisorischer Patch beseitigt ein Leck nicht vollständig. Der Bediener muss deshalb zwischen unmittelbarem Schiffserhalt und Wiedergewinnung von Kampffähigkeit abwägen. Mehrere Schäden können dieselben Zugangswege und Ressourcen beanspruchen. Autocrew kann Reparaturen konservativ zuteilen, bleibt aber an dieselbe Zustands- und Ressourcenwelt gebunden.')

add('SCHADEN & RÜCKKOPPLUNG', 'Schadenslage sichtbar machen', 'Repository-Abbildung: funktionale Zonen, Flutung, Feuer und Reparaturtrupps',
    [('docs/screenshots/de-damage-control-alert.png', 'Das Schema ist fiktiv und dient der Zuordnung von Systemfolgen.')],
    'Entscheidend ist nicht nur der Schadensprozentsatz, sondern die betroffene Fähigkeit und ihre Rückwirkung.',
    ['src/ui/stations_view.py', 'src/ship/damage.py', 'docs/screenshots/de-damage-control-alert.png'],
    'Die Abbildung bietet einen Anlass, den Zusammenhang zwischen Zonen und Stationen durchzugehen. Ein Maschinenschaden verringert Beweglichkeit; ein Sonarschaden schwächt Aufklärung; Flugdeckschaden betrifft eigene Luftmittel und im aktuellen Modell auch achtere Steuerfunktionen. Links- oder rechtsseitige Flutung kann zusätzlich Schlagseite erzeugen. Die Karte visualisiert ein funktionales Modell und darf nicht als maßstäblicher realer Fregattenplan gelesen werden.', 'screen')

add('BEDIENUNG & KOOPERATION', 'Neun Stationen bedienen denselben Zustand', 'Stationsgrenzen sind Aufgabenverteilungen, keine isolierten Simulationen',
    [('Brücke / Maschine', 'Kurs, Fahrt, Navigation, Energie', 'Wirkt auf Lärm, Geometrie und Deck'),
     ('Sonar / ELOKA / Funk', 'Messen, analysieren, melden', 'Liefert Kontakte und Evidenz'),
     ('OPZ / Waffen', 'Lage, Freigabe, Wirkmittel', 'Verbraucht frische Beobachtungen'),
     ('Helikopter', 'Verlegung, Dippen, Bojen, Torpedo', 'Erweitert Beobachter- und Startorte'),
     ('Schadensabwehr', 'Trupps, Feuer, Flutung, Stabilität', 'Erhält die übrigen Fähigkeiten')],
    'Ein Stationswechsel ändert den Blick und die Eingabezuständigkeit, nicht die zugrunde liegende Welt.',
    ['src/core/station.py', 'src/core/commands.py', 'src/core/game.py', 'src/core/help.py'],
    'Lokale Eingaben haben eine eindeutige Zuständigkeit. Menüs, Editor, Verwaltungsansichten und numerische Eingaben werden vor den normalen Stationsaktionen behandelt. Gehaltene Steuerungen werden bei relevanten Wechseln gelöscht. Numerische Eingaben können absichtlich bei laufender Simulation erfolgen. Kartenbedienung gilt nur dort, wo eine Karte sichtbar ist. Dieselbe klare Zuständigkeit wird für Remote Crew benötigt: Ein Browser soll nicht versehentlich zugleich eine andere Station oder den Host verwalten können.', 'table')

add('BEDIENUNG & KOOPERATION', 'Autocrew automatisiert vorhandene Arbeitsplätze', 'Begrenzte Regeln auf Basis veröffentlichter Beobachtungen',
    [('Eigene Takte', 'Stationen entscheiden alle 0,5 bis 5 s; Reihenfolge und nächste Fälligkeit sind festgelegt.'),
     ('Gleiche Grenzen', 'Schaden und fehlende Beobachtungen begrenzen Aktionen; keine allgemeine verborgene Zielkenntnis.'),
     ('Kooperation', 'Eine fernbelegte Station setzt ihre Autocrew aus; andere Stationen können weiter automatisiert sein.')],
    'Autocrew nutzt das zuletzt veröffentlichte Lagebild und setzt Befehle für folgende Simulationsschritte.',
    ['src/core/autocrew.py', 'src/core/game.py', 'tests/test_autocrew.py'],
    'Die Automatisierung ist ein Bündel stationsbezogener Regeln, kein lernendes Modell. Die Brücke kann beobachteten Bedrohungen ausweichen und Navigationsgefahren berücksichtigen; Sonar kann unterstützende Mess-/TMA-Aktionen ausführen; Schadensabwehr verteilt Trupps. Der Zustand enthält aktivierte Stationen, nächste Fälligkeiten und letzte Aktionen und wird für die Fortsetzung gespeichert. Die Ausführung erfolgt am Ende eines Substeps, damit neue Beobachtungen bereits verfügbar sind. Dadurch besteht eine klare Verzögerung zwischen Wahrnehmung und Wirkung des nächsten Stellbefehls.')

add('BEDIENUNG & KOOPERATION', 'Remote Crew: Der Host bleibt die Autorität', 'Browser übertragen Bedienabsichten und erhalten erlaubte Projektionen',
    [('Browser', 'Stationsbedienung erzeugt einen versionierten, begrenzten Befehl.'),
     ('Transport', 'Authentifizierung und Schema prüfen; abgetrennte Nachricht in eine begrenzte Queue legen.'),
     ('Hauptthread', 'Rolle, Lease, Epoche, Revision und Bereitschaft erneut prüfen; in stabiler Reihenfolge anwenden.'),
     ('Rückkanal', 'Beobachtungsbasierter Zustand, Ereignisse und freigegebenes Audio gehen an berechtigte Clients.')],
    'Protokoll v2, höchstens 12 Sitzungen und exklusive Stationsbelegung begrenzen den gemeinsamen Spielraum.',
    ['src/commander/server.py', 'src/commander/local.py', 'src/commander/bridge.py', 'src/commander/projections.py'],
    'Der Netzwerkthread führt keine Physik und keine beliebigen Game-Methoden aus. Nachrichten werden streng validiert und im Hauptthread erneut geprüft. Stationsreihenfolge und pro Client geordnete Befehle bestimmen die Anwendung; Sequenz- und Kontextprüfungen verhindern Doppelanwendung oder alte Befehle nach einem Weltwechsel. Der Host vergibt Rechte, inklusive gesonderter direkter Feuerfreigabe. Solo Crew und der optionale Web-Host erweitern die ausdrücklich freigegebene Host-Bedienoberfläche, ohne eine zweite Simulationsautorität zu erzeugen. Sitzungen und Zugangsdaten sind keine Bestandteile von Spielständen.', 'flow')

add('BEDIENUNG & KOOPERATION', 'Lokale und Browseransicht teilen Beobachtungen', 'Repository-Abbildung einer deutschen Remote-Crew-OPZ',
    [('docs/screenshots/commander-v2-de-opz-desktop.png', 'Mehrere Bediener arbeiten an derselben autoritativen Einsatzlage.')],
    'Netzwerkverzögerung kann Anzeigen altern lassen; daher wird beim Befehl die aktuelle Gültigkeit nochmals geprüft.',
    ['src/commander/projections.py', 'data/commander/app.js', 'docs/screenshots/commander-v2-de-opz-desktop.png'],
    'Der Browser rendert eine zulässige, abgetrennte Projektion. Er erhält keine frei lesbaren Simulationsobjekte und berechnet keine eigene authoritative Welt. Stationsrechte beschränken sowohl Aktionen als auch sichtbare Daten. Lokale, direkte und fernbediente Befehle treffen sich im Game-seitigen Handlungspfad. Die im Repository vorhandenen Screenshots sind Ansichtsbeispiele, keine aktuelle Latenz- oder Lastmessung. Eine gute Netzwerkdarstellung muss veraltete oder widerrufene Zustände für den Bediener erkennbar machen.', 'screen')

add('DATEN & FORTSETZUNG', 'Live-Verkehr ist eine externe Eingangsgröße', 'Optionales AIS/ADS-B ergänzt die erzeugte Welt',
    [('Import', 'AISStream-/OpenSky-Clients liefern Meldungen; Koordinaten werden in den realen Sektor projiziert.'),
     ('Integration', 'Der Hauptthread erzeugt/aktualisiert Einheiten; Relevanzradius, Alter und feste Obergrenzen begrenzen Last.'),
     ('Folge', 'Feedfolge und Wanduhr beeinflussen diesen Modus. Derselbe Seed allein reicht hier nicht zur Wiederholung.')],
    'Aktuelle Obergrenzen: 60 Live-Schiffe und 40 Live-Flugzeuge; Feedpflege läuft vor dem Physik-Pausegate.',
    ['src/network/live_traffic.py', 'src/network/ais_client.py', 'src/network/adsb_client.py', 'src/air/live_aircraft.py'],
    'Das ist eine wesentliche Grenze der Determinismusaussage. Der seedbasierte Kern kann bei identischen Eingaben fortgesetzt werden; ein externer Livestream ist selbst eine zusätzliche Eingabe mit zeitabhängigen Ankünften. Live-Flugzeuge werden im Pump anhand von time.time weitergeführt, Meldungen nach Wandzeit verworfen. Live-Schiffspositionsupdates sind bewusst gedrosselt. Die Pumpreihenfolge bedeutet, dass Netzwerkpflege auch während einer Kernsimulationspause stattfinden kann. Ein gespeicherter Weltzustand ist deshalb keine vollständige Aufzeichnung aller zukünftigen Live-Meldungen.')

add('DATEN & FORTSETZUNG', 'Determinismus braucht mehr als denselben Seed', 'Zustand, Reihenfolge und Eingabefolge müssen übereinstimmen',
    [('Zufall', 'Lokale Random-Streams mit gespeichertem Zustand; neue Physik nutzt stateless detrand-Schlüssel.'),
     ('Reihenfolge', 'Stabile Iteration und Update-Reihenfolge; gleiche Eingaben und gleiche Zeitschrittfolge.'),
     ('Trennung', 'Rendering, Audiogerät und reine Caches dürfen den taktischen Zustand nicht verändern.')],
    'Reproduzierbar ist die definierte Simulation mit ihren Eingängen – nicht eine beliebig andere Ausführung.',
    ['src/core/detrand.py', 'src/core/save_schema.py', 'tests/test_determinism.py', 'tests/test_runtime_continuation.py'],
    'Ein klassischer Zufallsstrom hängt von allen vorherigen Ziehungen ab. Neue Ziehungen können deshalb bestehende Abläufe verschieben. Das counterbasierte detrand bildet Schlüssel wie Seed, Tag, Entität und Tick direkt auf Zufallswerte ab. Speicherfortsetzung muss zusätzlich IDs, Sensorzyklen, ausstehende Ereignisse und Waffenphasen erhalten. Unterschiedliche dt-Sequenzen sind nicht pauschal identisch; einzelne Solver sind besonders substep-stabil, die gesamte Ablaufsteuerung hat jedoch diskrete Ereignisse. Live-Feeds benötigen für eine echte Wiederholung auch identische externe Meldungen.')

add('DATEN & FORTSETZUNG', 'Save v12 hält einen fortsetzbaren Weltzustand', 'Laden ist ein validierter Zustandswechsel',
    [('Was gespeichert wird', 'Welt/Küste, Entitäten, Katalogsnapshots, RNG, IDs, Sensoren, Waffen und ausstehende Ereignisse.'),
     ('Wie geschrieben wird', 'Temporäre Nachbardatei → flush/fsync → atomarer Austausch des Speicherplatzes.'),
     ('Wie geladen wird', 'Exaktes Schema und Werte prüfen → Kandidatenzustand herstellen → erst dann übernehmen.')],
    'Alte, neue oder unvollständige Save-Schemata werden abgelehnt; Netzwerkzugänge und Sitzungen bleiben transient.',
    ['src/core/game.py', 'src/core/save_schema.py', 'src/core/version.py', 'tests/test_save_v12_hygiene.py'],
    'Ein Spielstand ist mehr als eine Liste sichtbarer Einheiten. Er muss auch einen schon gesendeten Ping, eine laufende Waffenladung, eine gestartete Reparatur oder einen künftigen Sensorzyklus korrekt fortsetzen. Die Weltgeometrie und runtime-relevanten Katalogdaten sind enthalten, um Veränderungen an Generatoren oder installierten Katalogen nicht stillschweigend in einen alten Einsatz einzuschleusen. Validierung und Wiederherstellung erfolgen transaktional. Nicht jeder Darstellungszustand ist persistent: Empfängersignalpuffer haben einen absichtlichen Warm-up, OPZ-Fusionsarbeitsflächen sind transient.')

add('DATEN & FORTSETZUNG', 'Mission und Editor: Modellumfang ist nicht Laufzeitumfang', 'Die Startbrücke akzeptiert nur explizit unterstützte Definitionen',
    [('Mission', 'Szenario/Seed und Schwierigkeit bestimmen Kräfte, Zeitlimit und Zielmodus sink oder survive.'),
     ('Eigene Mission', 'Fester 500-NM-Raum, klares Wetter, unterstützte eingebaute U-Boot-/Oberflächenprofile.'),
     ('Bewusste Grenzen', 'Keine Laufzeit für Editor-Events, Zufallsgruppen oder eigene Einheitenprofile; ungeeignete Definitionen werden abgewiesen.')],
    'Editorvalidierung prüft eine Definition; erst die Laufzeitprüfung entscheidet, ob sie gespielt werden kann.',
    ['src/core/mission.py', 'src/core/mission_definition.py', 'src/core/game.py', 'src/data/user_content.py'],
    'Eingebaute Szenarien und die freie Jagd verwenden Mission als Spawn- und Zielplan. Die Missionsprüfung wertet den tatsächlichen Einsatzverlauf aus, einschließlich eigener Verluste, Vorfällen und Zeitlimit. Für editorverfasste Missionen existiert eine engere Runtime-Brücke: unterstützte eingebaute Plattformprofile, keine Events oder Zufallsgruppen und nur sink/survive. Im aktuellen start_custom_mission werden die Ziel-IDs für sink gegen die platzierten feindlichen U-Boote geprüft; die Seite wird beim Erzeugen einer Einheit übernommen. Benutzerinhalte werden mit strikten Typ-, Pfad- und Referenzprüfungen behandelt. Gültige benutzerdefinierte Einheiten haben aktuell keine Simulationseffekt-Zusage.')

add('AUSGABE & TECHNIK', 'Audio macht die akustische Lage hörbar', 'Synthese und Analyse bleiben von der Geräteverfügbarkeit getrennt',
    [('Erzeugung', 'Beobachtete Quellen → tonale Linien und kontinuierliches Bandrauschen → Beam-Mischung.'),
     ('Auswertung', '250-ms-Blöcke bei 4096 Hz; FFT, Breitbandverlauf und Hüllkurvenanalyse liefern Anzeigen.'),
     ('Wiedergabe', 'Hörmodus, Filter, Lautstärke und begrenzte Puffer; fehlendes Gerät führt zu Stille.')],
    'Im aktuellen Game-Pfad ist Sonarhören opt-in; es gibt keine dauernde Eigenschiff-Geräuschkulisse.',
    ['src/audio/receiver.py', 'src/audio/hydroacoustics.py', 'src/audio/engine.py', 'src/core/game.py'],
    'Die technischen Audiomodule bieten mehr Bausteine, als der aktuelle Game-Pfad ständig abspielt. _update_audio streamt das aktiv gewählte Sonar- oder Helikopterhören. Die Referenzsignale sind synthetisch und sollen Unterscheidungsaufgaben ermöglichen, keine realen geheimen Plattformaufnahmen reproduzieren. Der Empfänger hält Phasen- und Filterzustände über Blöcke. Wiedergabe und Browsertransport haben begrenzte Queues und behandeln Unterbrechungen explizit. Der tactical state darf nicht davon abhängen, ob ein Audiogerät erfolgreich initialisiert wurde.')

add('AUSGABE & TECHNIK', 'Technische Abhängigkeiten bleiben überschaubar', 'Laufzeit, Ressourcen und Präsentationswerkzeug sind getrennt',
    [('Python ≥ 3.11', 'Standardbibliothek', 'Zustände, JSON, Threads, HTTP, Dateisystem'),
     ('Pygame ≥ 2.6', 'Spiel-Laufzeit', 'Fenster, Eingaben, Zeichnen und Mixer'),
     ('NumPy ≥ 2.0', 'Numerische Berechnung', 'Signalverarbeitung und vektorisierte Modelle'),
     ('websockets ≥ 12', 'Netzwerkabhängigkeit', 'Unterstützung des Live-Datenzugangs'),
     ('Paketressourcen', 'JSON / Gzip / Browserdateien', 'Kataloge, Karten, Loadouts, Sprachen und UI')],
    'python-pptx, Pillow und ReportLab erzeugen diese Präsentation; sie sind keine neuen Spielabhängigkeiten.',
    ['pyproject.toml', 'requirements.txt', 'MANIFEST.in', 'src/network/ais_client.py'],
    'Die deklarierten Laufzeitabhängigkeiten stammen aus pyproject.toml. Viele andere Funktionen wie HTTP-Server, Speichertransaktionen, Hashes und Threads beruhen auf der Python-Standardbibliothek. Die Browseroberfläche liegt als paketierte HTML-, CSS- und JavaScript-Dateien vor. Kontaktanalysebilder sind inzwischen ebenfalls Paketressourcen; ältere pauschale Aussagen über keinerlei Bilddateien wären daher zu eng. Die Dokumenterzeugung installiert nur Werkzeuge in der vorhandenen venv; Projektmetadaten und Anforderungen des Spiels werden dadurch nicht verändert.', 'table')

add('AUSGABE & TECHNIK', 'Leistung wird durch Begrenzung gewonnen', 'uConsole: 1280 × 720 als Zielarbeitsfläche, 60 FPS als Zielrate',
    [('Rechenarbeit dosieren', 'Sensorintervalle, räumliche Vorauswahl, vektorisierte Signal-/Strahlrechnung und begrenzte Caches.'),
     ('Zustand begrenzen', 'Tracks, Historien, Wolken, Waffen, Netzwerkqueues und Live-Einheiten haben Obergrenzen.'),
     ('Ausgabe entkoppeln', 'Aspektgerechtes Letterboxing, beobachtungsbasierte Projektionen, begrenzte Audio-/Webarbeit.')],
    'Zielwerte sind keine Hardwaremessung; Thermik, Latenz und maximale Clientlast müssen auf dem Gerät geprüft werden.',
    ['src/core/config.py', 'src/core/game.py', 'src/sonar/raytrace.py', 'src/commander/server.py'],
    'Die reine Desktop- oder Headless-Ausführung belegt nicht, dass die volle Last auf der uConsole stabil bleibt. Besonders teuer sind viele Kandidaten pro Sensor, neue Raytrace-Tabellen und Signalverarbeitung. Darum werden Metadaten und reine Tabellen gecacht, abseits liegende Kandidaten früh verworfen und Publikationstakte begrenzt. Die Obergrenze von zwölf Sitzungen ist ein Protokoll-/Ressourcenlimit, kein gemessener Leistungsnachweis für jeden Clientmix. Der Code besitzt optionale Laufzeitdiagnostik für Physik-, Audio-, Commander- und Zeichenzeiten.')

add('ZUSAMMENSPIEL', 'Beispiel: Ein leiser Kontakt wird bekämpfbar', 'Illustrativer Ablauf aus implementierten Bausteinen',
    [('Suche', 'Fahrt reduzieren → weniger Eigenlärm → passiver Kontakt und Peilungshistorie.'),
     ('Geometrie', 'Manöver oder abgesetzter Sensor → neue Beobachterlage → belastbarere TMA.'),
     ('Freigabe', 'Kontakt bewerten → Fixfrische, ROE, Vorräte und Hülle prüfen → Startdatum setzen.'),
     ('Gegenwirkung', 'Gegner bemerkt Gefahr → Ausweichen/Täuschkörper → Suchkopf entscheidet → Schaden verändert Lage.')],
    'Jeder Übergang benötigt ein gültiges Ergebnis des vorherigen Systems; keiner garantiert den nächsten Erfolg.',
    ['src/core/game.py', 'src/sonar/tma.py', 'src/enemies/sub.py', 'src/weapons/torpedo.py'],
    'Dieses Beispiel ist keine aufgezeichnete Messreihe und verspricht keine festen Zeit- oder Reichweitenwerte. Es führt die zuvor beschriebenen Abhängigkeiten zusammen. Weniger Fahrt kann die Hörlage verbessern, aber die Manövergeometrie langsamer aufbauen. Ein Helikopter kann helfen, benötigt jedoch Flugwetter, Deck und Treibstoff. Eine TMA-Lösung kann bis zum Waffenstart wieder veralten. Nach dem Start kann der Gegner den Ausstoß oder die laufende Waffe bemerken. Wirkung und Folgelage hängen daher von mehreren unabhängigen begrenzenden Modellen ab.', 'flow')

add('ZUSAMMENSPIEL', 'Beispiel: Ein Wetterwechsel trifft die Einsatzkette', 'Eine Ursache, mehrere gleichzeitig veränderte Entscheidungen',
    [('Aufklärung', 'Mehr Wind/Regen → akustischer Hintergrund und Radarstörung ändern sich; Sicht sinkt.'),
     ('Plattform', 'Mehr Wellen → Widerstand und Deckbewegung wachsen; Flugfenster können sich schließen.'),
     ('Reaktion', 'Kurs/Fahrt oder Sensorwahl ändern → eigene Signatur, Geometrie und Ressourcen verändern sich erneut.')],
    'Umwelt ist eine gemeinsame Eingabe, die taktische Zielkonflikte erzeugt.',
    ['src/world/world.py', 'src/physics/ship_dynamics.py', 'src/core/game.py', 'src/sonar/equation.py'],
    'Auch dieses Beispiel beschreibt eine Kausalkette und keine garantierte monotone Veränderung jedes Messwertes. Ein verändertes Schallprofil kann manche Wege verschlechtern und andere verbessern. Die Flugentscheidung hängt zusätzlich von Windrichtung relativ zum Schiff und dem momentanen Deckzustand ab. Wenn die Fregatte zur Deckberuhigung den Kurs ändert, verändert das zugleich die Sonargeometrie und eventuell die Drahtbelastung einer laufenden Waffe. Genau solche Querverbindungen machen die Stationsaufteilung sinnvoll.')

add('QUALITÄT & GRENZEN', 'Tests prüfen verschiedene Arten von Richtigkeit', 'Verhalten, Fortsetzung, Beobachtungsgrenze und Spielkalibrierung',
    [('Physik & Wirkung', 'Bewegung, Sonargleichung, Strahlen, Waffen, Grundkontakt und Schadensstabilität.'),
     ('Systemverträge', 'Determinismus, Save-Roundtrip, Frische, Rollen, genau-einmal-Kommandos und Ressourcen.'),
     ('Bedienung & Balance', 'Layout, Sprachen, Browserzustände; 77 Referenzmetriken plus begründete Abweichungen.')],
    'Kalibrierung auf Spielreferenzen ist ein Balancenachweis innerhalb der Tests, keine marinefachliche Validierung.',
    ['tests/test_calibration.py', 'tests/calibration/golden.json', 'tests/calibration/deviations.json', 'tools/calibrate.py'],
    'Der Kalibrierungsharness vergleicht modellierte Ergebnisse mit gespeicherten Referenzen aus 1.0.0. Änderungen wie andere Mehrwegeausbreitung dürfen bewusst abweichen, wenn sie als bekannte Metrik mit Grund dokumentiert sind. Weitere Tests prüfen Verträge, die eine reine Physikformel nicht abdeckt: Keine versteckte Position in der Bedienansicht, keine alte Fernaktion nach Weltwechsel, kein Teilzustand nach ungültigem Laden. Für diese Dokumentationsarbeit werden gezielte Prüfungen protokolliert; sie sind kein Ersatz für einen vollständigen Release- oder Hardwaretest.')

add('QUALITÄT & GRENZEN', 'Welche Aussagen die Modelle nicht tragen', 'Vereinfachungen helfen, Ergebnisse richtig zu interpretieren',
    [('Physik', 'Spielkalibrierte Näherungen; synthetischer Meeresboden; begrenzte Strahlenzahl und vereinfachte Wassersäule.'),
     ('Wahrnehmung', 'Synthetische Signaturen; unsichere Klassifikation; manuelle OPZ-Fusion ohne allgemeine Peilungstriangulation.'),
     ('Betrieb', 'Editorumfang größer als Runtime; Live-Feeds extern zeitabhängig; kein vollständiger Luftkampf- oder Schiffsentwurfssimulator.')],
    'Die Stärke liegt in gekoppelten Entscheidungen und nachvollziehbaren Konsequenzen innerhalb des Spielmodells.',
    ['src/sonar/raytrace.py', 'src/audio/receiver.py', 'src/sensors/fusion.py', 'docs/simulation-gaps.md'],
    'Offene oder absichtlich ausgelassene Aspekte dürfen nicht aus vorhandenem Fachvokabular hinzugedacht werden. Eine Nomoto-Steuerung ist kein vollständiger CFD-Solver. Ein Temperaturprofil ist keine reale hydrographische Vermessung. Ein ESM-Kandidat ist keine gesicherte Plattformidentität. Der Implementierungsstand hat außerdem Vorrang vor historischen Designplänen und teilweise überholten Kommentaren. Besonders abweichend sind aktuell die feste Echtzeit, optionaler Live-Verkehr, abgeleitete Wetterinstrumente und der opt-in Audiopfad. Die Quelleninventur macht den dokumentierten Arbeitsstand nachvollziehbar.')

add('SCHLUSSBILD', 'Was greift ineinander?', 'Die fünf wichtigsten Abhängigkeiten für das Verständnis',
    [('Umwelt → Signal', 'Wetter, Tiefe und Boden verändern, was ein Sensor empfangen kann.'),
     ('Manöver → Wissen', 'Eigene Fahrt und Geometrie bestimmen Lärm, Peilungsqualität und TMA.'),
     ('Wissen → Wirkung', 'Trackfrische und Unsicherheit begrenzen einen sinnvollen Waffeneinsatz.'),
     ('Wirkung → Fähigkeit', 'Schaden verändert Mobilität, Sensoren, Waffen und Überlebensfähigkeit.'),
     ('Zustand → Zusammenarbeit', 'Host, Crew und Speicherung teilen dieselbe definierte Simulation.')],
    'U-Jagd verbindet physische Zustände mit begrenztem Wissen – daraus entstehen seine taktischen Entscheidungen.',
    ['src/core/game.py', 'src/core/save_schema.py', 'src/commander/projections.py'],
    'Zum Abschluss lassen sich alle Details auf diese fünf Ketten zurückführen. Keine Station optimiert allein: Eine Fahrtentscheidung verändert Sonar; ein Sonarmanöver verändert Draht und Deck; ein Treffer verändert die nächste Detektionschance. Der zentrale Integrationscode ordnet diese Wechselwirkungen zeitlich. Die Informationsgrenze verhindert, dass der Spieler schon weiß, was seine Sensoren erst herausfinden müssen. Gemeinsam mit begrenzten Vorräten, Reaktionszeiten und Schäden ergibt sich der taktische Charakter der Simulation.', 'summary')

add('ANHANG', 'Quellen und Lesewege', 'Die vollständigen Pfade stehen zusätzlich in den Foliennotizen und im Begleittext',
    [('Ablauf & Welt', 'src/core/game.py\nsrc/core/config.py\nsrc/world/\nsrc/physics/'),
     ('Information & Wirkung', 'src/sonar/ · src/sensors/\nsrc/ship/ · src/enemies/\nsrc/air/ · src/weapons/\nsrc/audio/'),
     ('Verträge & Nachweise', 'src/commander/ · src/network/\nsrc/data/ · data/\ntests/ · tools/calibrate.py\npyproject.toml')],
    'Beigelegt: Folien mit Sprechernotizen, Quelleninventur, Erzeugungsskript und Prüfprotokoll.',
    ['AGENTS.md', 'src/core/version.py', 'pyproject.toml'],
    'Priorität der Analyse: ausführbarer Code und fokussierte Tests, danach Paketressourcen und Metadaten, danach aktuelle Handbücher. Historische Designtexte wurden nur als Kontext verwendet. Es wurden keine privaten Referenz-PDFs verwendet. Die Architekturabbildungen wurden für diese Präsentation aus den beschriebenen Beziehungen neu erstellt. Bestehende Screenshots sind als Repository-Abbildungen gekennzeichnet. Das Quelleninventar enthält Dateihashes des lokalen Arbeitsstands; ein Commit-Hash allein würde die bereits vorhandenen lokalen Änderungen nicht hinreichend beschreiben.')


class Canvas:
    """Shared layout for editable PPTX objects and a deterministic visual preview."""
    def __init__(self, slide, pdf):
        self.slide = slide
        self.pdf = pdf
        self.image = Image.new('RGB', (W, H), BG)
        self.draw = ImageDraw.Draw(self.image)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BG[1:])
        pdf.setFillColor(HexColor(BG))
        pdf.rect(0, 0, W, H, stroke=0, fill=1)

    def box(self, x, y, w, h, fill=PANEL, radius=0):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x / 120), Inches(y / 120), Inches(w / 120), Inches(h / 120))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill[1:])
        shape.line.fill.background()
        self.pdf.setFillColor(HexColor(fill))
        self.pdf.rect(x, H-y-h, w, h, fill=1, stroke=0)
        if radius:
            self.draw.rounded_rectangle((x, y, x+w, y+h), radius, fill=fill)
        else:
            self.draw.rectangle((x, y, x+w, y+h), fill=fill)

    def line(self, x1, y1, x2, y2, color=MUTED):
        self.draw.line((x1,y1,x2,y2), fill=color, width=2)
        self.pdf.setStrokeColor(HexColor(color))
        self.pdf.setLineWidth(2)
        self.pdf.line(x1,H-y1,x2,H-y2)
        connector=self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1/120), Inches(y1/120),
            Inches(x2/120), Inches(y2/120))
        connector.line.color.rgb=RGBColor.from_string(color[1:])
        connector.line.width=Pt(1.2)

    def text(self, text, x, y, w, size=30, color=INK, bold=False, max_h=None):
        font = ImageFont.truetype(BOLD if bold else FONT, size)
        lines = []
        for paragraph in text.split('\n'):
            line = ''
            for word in paragraph.split():
                trial = f'{line} {word}'.strip()
                if font.getlength(trial) <= w:
                    line = trial
                else:
                    if not line:
                        raise ValueError(f'Unbreakable word: {word}')
                    lines.append(line)
                    line = word
            lines.append(line)
        step = size * 1.40
        height = len(lines) * step
        if max_h is not None and height > max_h:
            raise ValueError(f'Text overflow ({height:.1f}>{max_h}): {text[:100]}')
        if y + height > H - 12:
            raise ValueError(f'Out of canvas: {text[:100]}')
        for i, line in enumerate(lines):
            yy = y + i * step
            self.draw.text((x, yy), line, font=font, fill=color, anchor='lt')
            if not line:
                continue
            self.pdf.setFillColor(HexColor(color))
            self.pdf.setFont('DejaVu-Bold' if bold else 'DejaVu', size)
            self.pdf.drawString(x, H-yy-size*.76, line)
            tb = self.slide.shapes.add_textbox(Inches(x/120), Inches((yy-2)/120),
                                              Inches((w+12)/120), Inches(step/120))
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.text = line
            p.font.name = 'DejaVu Sans'
            p.font.size = Pt(size * .6)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color[1:])
        return height

    def picture(self, path, x, y, w, h):
        with Image.open(path) as im:
            im = im.convert('RGB')
            ratio = min(w / im.width, h / im.height)
            sw, sh = int(im.width*ratio), int(im.height*ratio)
            xx, yy = x + (w-sw)//2, y + (h-sh)//2
            self.image.paste(im.resize((sw, sh), Image.Resampling.LANCZOS), (xx, yy))
        self.slide.shapes.add_picture(str(path), Inches(xx/120), Inches(yy/120),
                                      width=Inches(sw/120), height=Inches(sh/120))
        self.pdf.drawImage(str(path), xx, H-yy-sh, sw, sh)


def render(prs, pdf, row, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    c = Canvas(slide, pdf)
    c.box(0, 0, 12, H, TEAL)
    c.text(row['section'], 65, 38, 1200, 18, TEAL, True)
    c.text(f'{number:02d} / {len(SLIDES):02d}', 1420, 38, 120, 18, MUTED)
    title_size = 44
    while ImageFont.truetype(BOLD, title_size).getlength(row['title']) > 1470:
        title_size -= 1
    c.text(row['title'], 65, 91, 1470, title_size, bold=True, max_h=70)
    c.text(row['subtitle'], 65, 166, 1470, 25, MUTED, max_h=72)
    cards = row['cards']
    kind = row['kind']
    if kind == 'cover':
        c.text('WELT → MESSUNG → ENTSCHEIDUNG → WIRKUNG', 65, 305, 1480,
               33, TEAL, True, max_h=100)
        for i, (title, body) in enumerate(cards):
            x = 65 + i*503
            c.box(x, 445, 477, 204)
            c.text(title, x+25, 477, 425, 42, GOLD, True, 65)
            c.text(body, x+25, 553, 425, 24, INK, max_h=85)
    elif kind in ('flow',):
        n = len(cards)
        width = (1470-(n-1)*32)/n
        for i, (title, body) in enumerate(cards):
            x = 65+i*(width+32)
            c.box(x, 275, width, 440)
            c.text(f'{i+1:02d}', x+23, 297, width-46, 46, TEAL, True)
            c.text(title, x+23, 377, width-46, 28, INK, True, 82)
            c.text(body, x+23, 465, width-46, 24, MUTED, max_h=240)
            if i < n-1:
                c.text('→', x+width+2, 417, 30, 27, GOLD)
    elif kind == 'table':
        x = [65, 365, 910]
        widths = [280, 525, 625]
        for j, title in enumerate(('BEREICH', 'MODELL / INHALT', 'FOLGE / VERWENDUNG')):
            c.text(title, x[j]+15, 250, widths[j]-25, 19, TEAL, True)
        rh = 83 if len(cards) == 5 else 101
        for i, cells in enumerate(cards):
            y = 295+i*rh
            c.box(65, y, 1470, rh-7)
            for j, cell in enumerate(cells):
                c.text(cell, x[j]+15, y+14, widths[j]-28, 23,
                       INK if j==0 else MUTED, j==0, rh-14)
    elif kind == 'screen':
        path, caption = cards[0]
        c.picture(ROOT/path, 65, 238, 1120, 492)
        c.box(1220, 257, 315, 427)
        c.text('ANSICHTSBEISPIEL', 1240, 285, 275, 19, TEAL, True)
        c.text(caption, 1240, 345, 275, 27, INK, max_h=290)
    elif kind == 'architecture':
        for left, hub, edge in ((525,555,585),(1075,1045,1015)):
            for yy, target_y in ((321,440),(483,483),(645,525)):
                c.line(left,yy,hub,yy)
                c.line(hub,yy,hub,target_y)
                c.line(hub,target_y,edge,target_y)
        positions = [(65,255,460,132),(65,417,460,132),(65,579,460,132),
                     (1075,255,460,132),(585,395,430,175),
                     (1075,417,460,132),(1075,579,460,132)]
        for i, ((title, body), (x,y,w,h)) in enumerate(zip(cards, positions)):
            c.box(x,y,w,h, '#22544F' if i==4 else PANEL)
            c.text(title,x+22,y+19,w-44,27,TEAL if i==4 else INK,True,78)
            c.text(body,x+22,y+65,w-44,22,MUTED,max_h=h-68)
        c.text('Orchestriert den\nZustandsübergang', 617,591,380,23,MUTED)
    elif kind == 'summary':
        for i,(title,body) in enumerate(cards):
            y=260+i*87
            c.text(f'{i+1:02d}',65,y,80,32,TEAL,True)
            c.text(title,160,y,425,28,INK,True,80)
            c.text(body,620,y,910,26,MUTED,max_h=80)
    else:
        n = len(cards)
        width = (1470-(n-1)*27)/n
        for i,(title,body) in enumerate(cards):
            x=65+i*(width+27)
            c.box(x,267,width,422)
            c.box(x,267,width,5,TEAL if i%2==0 else GOLD)
            title_h=c.text(title,x+25,300,width-50,30,INK,True,130)
            body_y=max(391,320+title_h)
            c.text(body,x+25,body_y,width-50,27,MUTED,max_h=665-body_y)
    c.box(65,751,1470,3,TEAL)
    c.text(row['takeaway'],65,779,1460,25,INK,False,76)
    citation=' · '.join(row['sources'][:3])
    if len(citation)>165:
        citation=citation[:162]+'…'
    c.text(citation,65,868,1470,14,MUTED,max_h=25)
    slide.notes_slide.notes_text_frame.text = (
        f"Folie {number}: {row['title']}\n\n{row['notes']}\n\nQUELLEN\n"+
        '\n'.join(row['sources']))
    return c.image


def inventory():
    rows=[]
    counts={}
    for p in sorted((ROOT/'src').rglob('*.py')):
        rel=p.relative_to(ROOT).as_posix()
        code=p.read_text(encoding='utf-8')
        tree=ast.parse(code)
        imports=set()
        defs=[]
        for node in ast.walk(tree):
            if isinstance(node,ast.ImportFrom) and node.module and node.module.startswith('src'):
                imports.add(node.module)
            elif isinstance(node,ast.Import):
                imports.update(n.name for n in node.names if n.name.startswith('src'))
            elif isinstance(node,(ast.ClassDef,ast.FunctionDef,ast.AsyncFunctionDef)):
                defs.append(f'{node.name}:{node.lineno}')
        group=rel.split('/')[1]
        counts[group]=counts.get(group,0)+1
        rows.append(dict(path=rel,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),
                         lines=len(code.splitlines()),docstring=ast.get_docstring(tree),
                         internal_imports=sorted(imports),symbols=defs))
    resources=[]
    patterns=('data/**/*.json','data/**/*.json.gz','data/manual/*.de.md',
              'data/commander/*.js','data/commander/*.css','data/commander/*.html',
              'tests/test_*.py','tests/calibration/*.json','tools/*.py')
    for p in sorted({p for pattern in patterns for p in ROOT.glob(pattern)}):
        resources.append(dict(path=p.relative_to(ROOT).as_posix(),
                              sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
    return dict(date='2026-09-24',version='1.1.0',
                base_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
                working_tree=subprocess.check_output(['git','status','--short'],cwd=ROOT,text=True),
                scope='AST-Inventur aller src-Pythonmodule; fachliche Vertiefung gemäß Folienquellen. Keine Behauptung einer vollständigen Zeilenprüfung.',
                counts=counts,modules=rows,resources=resources)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    prs=Presentation()
    prs.slide_width=Inches(W/120)
    prs.slide_height=Inches(H/120)
    prs.core_properties.title='U-Jagd – Simulation und Abhängigkeiten'
    prs.core_properties.subject='Architektur und gekoppelte Simulationsmodelle im lokalen Arbeitsstand 1.1.0'
    prs.core_properties.author='U-Jagd · technische Dokumentation'
    pdfmetrics.registerFont(TTFont('DejaVu',FONT))
    pdfmetrics.registerFont(TTFont('DejaVu-Bold',BOLD))
    pdf=PDFCanvas(str(OUT/'U-Jagd-Simulation.de.pdf'),pagesize=(W*.6,H*.6))
    pdf.setTitle(prs.core_properties.title)
    pdf.setAuthor(prs.core_properties.author)
    previews=[]
    for i,row in enumerate(SLIDES,1):
        print(f'{i:02d} {row["title"]}',flush=True)
        pdf.saveState()
        pdf.scale(.6,.6)
        previews.append(render(prs,pdf,row,i))
        pdf.restoreState()
        pdf.showPage()
    prs.save(OUT/'U-Jagd-Simulation.de.pptx')
    pdf.save()
    thumb_w,thumb_h=400,225
    cols=4
    sheet=Image.new('RGB',(cols*thumb_w,((len(previews)+cols-1)//cols)*thumb_h),BG)
    for i,im in enumerate(previews):
        sheet.paste(im.resize((thumb_w,thumb_h),Image.Resampling.LANCZOS),
                    ((i%cols)*thumb_w,(i//cols)*thumb_h))
    sheet.save(OUT/'Folienuebersicht.jpg',quality=90)
    # A few full-size views support visual inspection without retaining 50+ PNGs.
    for index in (0,3,7,16,23,39,len(previews)-1):
        previews[index].save(OUT/f'vorschau-{index+1:02d}.png')
    md=['# U-Jagd: Simulation und Abhängigkeiten',
        '\nStand: lokaler Arbeitsbaum vom 24.09.2026, Anwendung 1.1.0.\n',
        'Bearbeitbare Folien: `U-Jagd-Simulation.de.pptx`; visuelle Leseversion: `U-Jagd-Simulation.de.pdf`.\n']
    for i,row in enumerate(SLIDES,1):
        md += [f'\n## {i:02d} · {row["title"]}\n',row['subtitle']+'\n']
        for card in row['cards']:
            md.append('- '+' — '.join(card).replace('\n','; '))
        md += ['\n'+row['takeaway']+'\n','### Sprechernotizen\n',row['notes']+'\n',
               '### Quellen\n']
        for src in row['sources']:
            assert (ROOT/src).exists(),src
            md.append(f'- [{src}](../../{src})')
    (OUT/'Sprechernotizen-und-Quellen.de.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    inv=inventory()
    (OUT/'Quelleninventar.json').write_text(json.dumps(inv,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'Built {len(SLIDES)} slides. Inventoried {len(inv["modules"])} Python modules.')


if __name__=='__main__':
    main()
